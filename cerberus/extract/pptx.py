"""PPTX extractor: a slide deck -> block IR.

A slide is a pile of positioned shapes stored in drawing (z) order, with no
reading order and usually no grouping. Three steps recover the structure:

1. Collect every shape with its position on the slide, unpacking group shapes.
2. Nest: a text-less shape that encloses other shapes (a card background, a
   header band) becomes a GROUP, so a name and its email and phone share a parent.
3. Order siblings with a recursive XY-cut: split into rows where a horizontal
   gap runs right across, otherwise into columns, and repeat inside each piece.

Content python-pptx cannot read (charts, SmartArt, embedded objects) quarantines
the slide: fail closed rather than pass text nobody inspected.
"""

from __future__ import annotations

from dataclasses import dataclass, field
from pathlib import Path
from statistics import median

import pptx
from pptx.shapes.graphfrm import GraphicFrame
from pptx.shapes.group import GroupShape
from pptx.shapes.picture import Picture

from . import images
from .ir import BBox, Block, BlockType, Container, Document, Page, Unit, doc_id_from_bytes

# File properties that can name a person; they travel with the file invisibly.
_CORE_PROPERTIES = ("title", "subject", "author", "last_modified_by", "keywords", "comments", "category")

_IDENTITY = (1.0, 1.0, 0.0, 0.0)  # scale x, scale y, offset x, offset y


@dataclass
class _Node:
    kind: str  # "text", "table", "picture", "group", or "frame" (text-less shape)
    shape: object
    left: float
    top: float
    right: float
    bottom: float
    part: str  # "slide", "layout" or "master"
    children: list[_Node] = field(default_factory=list)

    @property
    def area(self) -> float:
        return (self.right - self.left) * (self.bottom - self.top)

    def encloses(self, other: _Node) -> bool:
        cx, cy = (other.left + other.right) / 2, (other.top + other.bottom) / 2
        return (
            self.area > other.area
            and self.left <= cx <= self.right
            and self.top <= cy <= self.bottom
        )


def extract_pptx(path: str | Path) -> Document:
    path = Path(path)
    doc = Document(
        doc_id_from_bytes(path.read_bytes()),
        path.name,
        "pptx",
        versions={"python-pptx": pptx.__version__},
    )
    prs = pptx.Presentation(str(path))

    for name in _CORE_PROPERTIES:
        value = getattr(prs.core_properties, name)
        if value:
            doc.add(BlockType.METADATA, value, native_ref={"core_property": name})

    reader = images.Reader()  # one per deck, so a logo on every slide is read once
    # Rows or columns overlapping by less than this still count as separate.
    tolerance = (prs.slide_width / 200, prs.slide_height / 200)

    for number, slide in enumerate(prs.slides, start=1):
        page = Page(number, prs.slide_width, prs.slide_height, Unit.EMU)
        doc.pages.append(page)
        slide_block = doc.add(BlockType.SLIDE, page=number, native_ref={"slide_id": slide.slide_id})
        first = len(doc.blocks)

        unsupported: list[str] = []
        layout = slide.slide_layout
        nodes = (
            # Text typed onto the master or layout shows on the slide but is not in it.
            _collect(layout.slide_master.shapes, _IDENTITY, "master", unsupported)
            + _collect(layout.shapes, _IDENTITY, "layout", unsupported)
            + _collect(slide.shapes, _IDENTITY, "slide", unsupported)
        )
        fonts: dict[str, float] = {}
        for node in _reading_order(_nest(nodes), tolerance):
            _emit(doc, node, slide_block.block_id, number, tolerance, fonts, reader)

        if slide.has_notes_slide and slide.notes_slide.notes_text_frame is not None:
            notes = _frame_text(slide.notes_slide.notes_text_frame)
            if notes.strip():
                doc.add(BlockType.NOTES, notes, parent_id=slide_block.block_id, page=number)

        content = doc.blocks[first:]
        title = _pick_title(slide, content, fonts)
        if title is not None:
            title.type = BlockType.HEADING
            path_ = (" ".join(title.text.split()),)
            for block in content:
                if block is not title:
                    block.section_path = path_

        if unsupported:
            page.quarantined = True
            page.reason = "unreadable content: " + ", ".join(sorted(set(unsupported)))

    errors = doc.validate()
    if errors:
        raise ValueError(f"{path.name}: extraction broke IR invariants: {errors[:5]}")
    return doc


def _collect(shapes, transform, part: str, unsupported: list[str]) -> list[_Node]:
    sx, sy, ox, oy = transform
    nodes = []
    for s in shapes:
        # Placeholders on a layout or master are prompts ("Click to add title"), not content.
        if part != "slide" and s.is_placeholder:
            continue
        left, top = ox + sx * (s.left or 0), oy + sy * (s.top or 0)
        node = _Node("frame", s, left, top, left + sx * (s.width or 0), top + sy * (s.height or 0), part)
        if isinstance(s, GroupShape):
            node.kind = "group"
            node.children = _nest(_collect(s.shapes, _group_transform(s, transform), part, unsupported))
            if not node.children:
                continue
        elif isinstance(s, GraphicFrame):
            if not s.has_table:
                unsupported.append("chart" if s.has_chart else "embedded object or diagram")
                continue
            node.kind = "table"
        elif isinstance(s, Picture):
            node.kind = "picture"
        elif s.has_text_frame and s.text_frame.text.strip():
            node.kind = "text"
        nodes.append(node)
    return nodes


def _group_transform(group: GroupShape, transform):
    """Children of a group are positioned in the group's own coordinate space."""
    sx, sy, ox, oy = transform
    xfrm = group._element.grpSpPr.xfrm
    if xfrm is None or xfrm.chOff is None or xfrm.chExt is None:
        return transform
    if not xfrm.chExt.cx or not xfrm.chExt.cy:
        return transform
    kx, ky = xfrm.ext.cx / xfrm.chExt.cx, xfrm.ext.cy / xfrm.chExt.cy
    return (
        sx * kx,
        sy * ky,
        ox + sx * (xfrm.off.x - xfrm.chOff.x * kx),
        oy + sy * (xfrm.off.y - xfrm.chOff.y * ky),
    )


def _nest(nodes: list[_Node]) -> list[_Node]:
    """Hang each node under the smallest text-less shape that encloses it.

    Smallest first, so a frame has received all its children before it is
    itself placed. Frames that end up enclosing nothing are decoration (lines,
    accent bars) and are dropped.
    """
    frames = sorted((n for n in nodes if n.kind == "frame"), key=lambda n: n.area)
    top_level = []
    for node in sorted(nodes, key=lambda n: n.area):
        if node.kind == "frame" and not node.children:
            continue
        parent = next((f for f in frames if f is not node and f.encloses(node)), None)
        (parent.children if parent else top_level).append(node)
    return top_level


def _reading_order(nodes: list[_Node], tolerance) -> list[_Node]:
    """Recursive XY-cut. Rows are tried before columns: people read across, then down."""
    if len(nodes) < 2:
        return nodes
    for lo, hi, tol in (("top", "bottom", tolerance[1]), ("left", "right", tolerance[0])):
        parts = _split(nodes, lo, hi, tol)
        if len(parts) > 1:
            return [n for part in parts for n in _reading_order(part, tolerance)]
    return sorted(nodes, key=lambda n: (n.top, n.left))  # overlapping shapes: no clean cut


def _split(nodes: list[_Node], lo: str, hi: str, tol: float) -> list[list[_Node]]:
    parts: list[list[_Node]] = []
    end = None
    for node in sorted(nodes, key=lambda n: (getattr(n, lo), getattr(n, hi))):
        if end is None or getattr(node, lo) >= end - tol:
            parts.append([])
            end = getattr(node, hi)
        parts[-1].append(node)
        end = max(end, getattr(node, hi))
    return parts


def _emit(doc: Document, node: _Node, parent_id: str, page: int, tolerance, fonts, reader) -> None:
    s = node.shape
    common = dict(
        parent_id=parent_id,
        page=page,
        bbox=BBox(node.left, node.top, node.right, node.bottom, Unit.EMU),
        native_ref={"part": node.part, "shape_id": s.shape_id},
    )

    if node.children:  # an explicit group, or a frame that encloses other shapes
        group = doc.add(BlockType.GROUP, **common)
        for child in _reading_order(node.children, tolerance):
            _emit(doc, child, group.block_id, page, tolerance, fonts, reader)

    elif node.kind == "text":
        block = doc.add(BlockType.TEXT_BOX, _frame_text(s.text_frame), **common)
        runs = [r for p in s.text_frame.paragraphs for r in p.runs]
        sizes = [r.font.size.pt for r in runs if r.font.size]
        if sizes:
            fonts[block.block_id] = max(sizes)
        # A link target (e.g. mailto:) is not visible on the slide but is in the file.
        for url in dict.fromkeys(r.hyperlink.address for r in runs if r.hyperlink.address):
            doc.add(
                BlockType.METADATA,
                url,
                parent_id=block.block_id,
                page=page,
                native_ref={**common["native_ref"], "hyperlink": True},
            )

    elif node.kind == "table":
        _emit_table(doc, node, common)

    elif node.kind == "picture":
        ref = dict(common["native_ref"])
        picture = None
        try:
            picture = s.image
            ref.update(sha1=picture.sha1, content_type=picture.content_type)
        except ValueError:  # linked, not embedded: no bytes in the file to OCR
            ref.update(linked=True)
        image = doc.add(
            BlockType.IMAGE_REGION, **{**common, "native_ref": ref}, container=Container.IMAGE
        )
        alt = s._element.xpath("./p:nvPicPr/p:cNvPr/@descr")
        if alt and alt[0].strip():
            doc.add(
                BlockType.METADATA,
                alt[0],
                parent_id=image.block_id,
                page=page,
                native_ref={**common["native_ref"], "alt_text": True},
            )
        if picture is not None:
            images.emit(doc, image, reader.read(picture.blob, picture.content_type))


def _emit_table(doc: Document, node: _Node, common: dict) -> None:
    table = node.shape.table
    block = doc.add(BlockType.TABLE, **common, container=Container.TABLE)

    # Cell boxes from the declared column widths and row heights, scaled to the frame.
    xs = _edges([c.width for c in table.columns], node.left, node.right)
    ys = _edges([r.height for r in table.rows], node.top, node.bottom)

    # Row 0 is taken as the header. The file's "header row" flag is not reliable:
    # tables styled by hand leave it off.
    has_header = len(table.rows) > 1
    headers: dict[int, str] = {}
    for r, row in enumerate(table.rows):
        for c, cell in enumerate(row.cells):
            if cell.is_spanned:  # swallowed by a merged cell; the origin cell carries the text
                continue
            width = cell.span_width if cell.is_merge_origin else 1
            height = cell.span_height if cell.is_merge_origin else 1
            cell_block = doc.add(
                BlockType.TABLE_CELL,
                _frame_text(cell.text_frame),
                container=Container.TABLE,
                parent_id=block.block_id,
                page=common["page"],
                bbox=BBox(xs[c], ys[r], xs[c + width], ys[r + height], Unit.EMU),
                native_ref=common["native_ref"],
                row=r,
                col=c,
                header_id=headers.get(c) if r else None,
            )
            if r == 0 and has_header:
                for covered in range(c, c + width):
                    headers[covered] = cell_block.block_id


def _edges(sizes: list[int], start: float, end: float) -> list[float]:
    total = sum(sizes) or 1
    edges, position = [start], 0
    for size in sizes:
        position += size
        edges.append(start + (end - start) * position / total)
    return edges


def _frame_text(text_frame) -> str:
    # python-pptx renders a soft line break as a vertical tab; swap it one-for-one
    # so character offsets still line up with the file.
    return "\n".join(p.text for p in text_frame.paragraphs).replace("\v", "\n")


def _pick_title(slide, content: list[Block], fonts: dict[str, float]) -> Block | None:
    """The title placeholder if the slide has one, else the largest text on the slide.

    Decks drawn with plain rectangles have no title placeholder. The largest
    text only counts if it is bigger than the slide's typical text; ties go to
    the first in reading order.
    """
    placeholder = slide.shapes.title
    if placeholder is not None:
        for block in content:
            ref = block.native_ref or {}
            if (
                block.type is BlockType.TEXT_BOX
                and ref.get("part") == "slide"
                and ref.get("shape_id") == placeholder.shape_id
            ):
                return block
    sized = [b for b in content if b.block_id in fonts]
    if len(sized) < 2:
        return None
    best = max(sized, key=lambda b: fonts[b.block_id])
    return best if fonts[best.block_id] > median(fonts.values()) else None
