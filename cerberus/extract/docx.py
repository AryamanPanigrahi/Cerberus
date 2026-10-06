"""DOCX extractor: a Word document -> block IR.

A .docx is a zip of XML parts. The body is a sequence of paragraphs and tables
in reading order, so no geometry is needed; the work is in reading everything:

- Walk the body XML in order. The library's `doc.paragraphs` skips tables,
  content controls and text boxes, and loses how they interleave.
- Structure comes from styles: a paragraph is a heading if its style carries
  an outline level, a list item if it carries numbering.
- Text that is in the file but not on the page is emitted too: file
  properties, headers and footers, notes and comments, link targets, picture
  alt text and tracked deletions.

A Word file stores no page numbers (pagination happens when it is displayed),
so location is the heading path plus the paragraph's index in its part.
"""

from __future__ import annotations

import hashlib
import re
from pathlib import Path

import docx
from docx.oxml import parse_xml

from . import images
from .ir import BlockType, Container, Document, doc_id_from_bytes

_W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
_R = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
_A = "http://schemas.openxmlformats.org/drawingml/2006/main"
_WP = "http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing"
_MC = "http://schemas.openxmlformats.org/markup-compatibility/2006"
_V = "urn:schemas-microsoft-com:vml"


def _w(tag: str) -> str:
    return f"{{{_W}}}{tag}"


_P, _TBL, _TR, _TC, _SDT_CONTENT = _w("p"), _w("tbl"), _w("tr"), _w("tc"), _w("sdtContent")
_VAL, _AUTHOR = _w("val"), _w("author")
_SKIP = {_w("pPr"), _w("rPr")}  # formatting only; pPr also holds tab-stop definitions
_GRAPHICS = {_w("drawing"), _w("pict"), _w("object")}
_DELETED = {_w("del"), _w("moveFrom")}
_INSERTED = {_w("ins"), _w("moveTo")}
_NOT_A_NOTE = {"separator", "continuationSeparator", "continuationNotice"}

_CORE_PROPERTIES = ("title", "subject", "author", "last_modified_by", "keywords", "comments", "category")
_APP_PROPERTIES = ("Company", "Manager")


def extract_docx(path: str | Path) -> Document:
    path = Path(path)
    doc = Document(
        doc_id_from_bytes(path.read_bytes()),
        path.name,
        "docx",
        versions={"python-docx": docx.__version__},
    )
    src = docx.Document(str(path))
    styles = _style_table(src.styles.element)

    reader = images.Reader()  # one per document, so a repeated picture is read once
    _emit_properties(doc, src)
    _Walker(doc, src.part, src.element.body, styles, reader).walk(src.element.body)

    related = sorted(src.part.rels.values(), key=lambda rel: rel.target_ref)
    for rel in related:
        kind = rel.reltype.rsplit("/", 1)[-1]
        if rel.is_external or kind not in ("header", "footer"):
            continue
        part = rel.target_part
        type_ = BlockType.HEADER if kind == "header" else BlockType.FOOTER
        _emit_container(doc, type_, part, part.element, part.element, styles, reader, {})

    for rel in related:
        kind = rel.reltype.rsplit("/", 1)[-1]
        if rel.is_external or kind not in ("footnotes", "endnotes", "comments"):
            continue
        part = rel.target_part
        root = getattr(part, "element", None)
        if root is None:  # the library has no class for this part; parse it ourselves
            root = parse_xml(part.blob)
        for note in root:
            if note.get(_w("type")) in _NOT_A_NOTE:
                continue
            ref = {"kind": kind, "id": note.get(_w("id"))}
            _emit_container(
                doc, BlockType.NOTES, part, root, note, styles, reader, ref, note.get(_AUTHOR)
            )

    errors = doc.validate()
    if errors:
        raise ValueError(f"{path.name}: extraction broke IR invariants: {errors[:5]}")
    return doc


def _emit_properties(doc: Document, src) -> None:
    """Fields saved with the file that can name a person or organisation."""
    for name in _CORE_PROPERTIES:
        value = getattr(src.core_properties, name)
        if value:
            doc.add(BlockType.METADATA, value, native_ref={"core_property": name})

    for part in sorted(src.part.package.iter_parts(), key=lambda p: str(p.partname)):
        name = str(part.partname)
        if name == "/docProps/app.xml":
            wanted = _APP_PROPERTIES
        elif re.fullmatch(r"/customXml/item\d+\.xml", name):
            wanted = None  # every field, but only of the cover-page schema below
        else:
            continue
        root = parse_xml(part.blob)
        if wanted is None and _local(root) != "CoverPageProperties":
            continue
        for leaf in root.iter():
            if not isinstance(leaf.tag, str) or len(leaf) or not (leaf.text or "").strip():
                continue
            if wanted is None or _local(leaf) in wanted:
                key = "cover_property" if wanted is None else "app_property"
                doc.add(BlockType.METADATA, leaf.text.strip(), native_ref={key: _local(leaf)})


def _emit_container(doc, type_, part, root, element, styles, reader, ref, author=None) -> None:
    """A header, footer or note: one parent block, dropped again if it holds nothing."""
    container = doc.add(type_, native_ref={"part": str(part.partname), **ref})
    before = len(doc.blocks)
    if author:
        doc.add(
            BlockType.METADATA,
            author,
            parent_id=container.block_id,
            native_ref={**container.native_ref, "comment_author": True},
        )
    _Walker(doc, part, root, styles, reader).walk(element, container.block_id)
    if len(doc.blocks) == before:
        doc.blocks.pop()


class _Walker:
    """Walks the block-level content of one part (body, a header, a note...)."""

    def __init__(self, doc: Document, part, root, styles: dict, reader: images.Reader):
        self.doc = doc
        self.reader = reader
        self.part = part
        self.name = str(part.partname)
        self.styles = styles
        # Position of every paragraph and table in this part: the write-back address.
        self.p_index = {p: i for i, p in enumerate(root.iter(_P))}
        self.tbl_index = {t: i for i, t in enumerate(root.iter(_TBL))}
        self.headings: list[tuple[int, object]] = []  # open headings as (level, block)

    def walk(self, container, parent_id: str | None = None, type_: BlockType | None = None) -> None:
        """parent_id=None means "under the current heading" (top level of the body)."""
        for element in _block_children(container):
            if element.tag == _P:
                self._paragraph(element, parent_id, type_)
            else:
                self._table(element, parent_id)

    def _section(self) -> tuple[str, ...]:
        return tuple(" ".join(block.text.split()) for _, block in self.headings)

    def _parent(self, parent_id: str | None) -> str | None:
        if parent_id is not None:
            return parent_id
        return self.headings[-1][1].block_id if self.headings else None

    def _paragraph(self, p, parent_id, type_) -> None:
        text, extras = self._inline(p)
        ref = {"part": self.name, "p": [self.p_index[p]]}
        owner = self._parent(parent_id)

        if text.strip():
            level, is_list = self._classify(p)
            if type_ is None and parent_id is None and level is not None:
                # A heading closes every open heading at its level or deeper.
                while self.headings and self.headings[-1][0] >= level:
                    self.headings.pop()
                block = self.doc.add(
                    BlockType.HEADING,
                    text,
                    parent_id=self._parent(None),
                    section_path=self._section(),
                    native_ref={**ref, "level": level},
                )
                self.headings.append((level, block))
            else:
                kind = type_ or (BlockType.LIST_ITEM if is_list else BlockType.PARAGRAPH)
                block = self.doc.add(
                    kind, text, parent_id=owner, section_path=self._section(), native_ref=ref
                )
            owner = block.block_id
        self._emit_extras(extras, owner, ref)

    def _table(self, tbl, parent_id) -> None:
        ref = {"part": self.name, "tbl": self.tbl_index[tbl]}
        table = self.doc.add(
            BlockType.TABLE,
            container=Container.TABLE,
            parent_id=self._parent(parent_id),
            section_path=self._section(),
            native_ref=ref,
        )
        rows = list(tbl.iterchildren(_TR))
        grid = tbl.find(_w("tblGrid"))
        columns = len(grid) if grid is not None else 0
        # Row 0 is taken as the header (the file's own header-row flag is rarely
        # set), except in one-column or one-row tables, which are layout boxes.
        has_header = len(rows) > 1 and columns > 1
        headers: dict[int, str] = {}

        for r, tr in enumerate(rows):
            col = _int(tr.find(f"{_w('trPr')}/{_w('gridBefore')}"), 0)
            for tc in tr.iterchildren(_TC):
                props = tc.find(_w("tcPr"))
                span = _int(props.find(_w("gridSpan")), 1) if props is not None else 1
                merge = props.find(_w("vMerge")) if props is not None else None
                continues = merge is not None and merge.get(_VAL) in (None, "continue")
                cell = self._cell(tc, table, r, col, headers.get(col) if r else None, continues)
                if cell is not None and r == 0 and has_header:
                    for covered in range(col, col + span):
                        headers[covered] = cell.block_id
                col += span

    def _cell(self, tc, table, row, col, header_id, continues):
        texts, indexes, extras_list, nested = [], [], [], []
        for element in _block_children(tc):
            if element.tag == _TBL:
                nested.append(element)
                continue
            text, extras = self._inline(element)
            extras_list.append((extras, self.p_index[element]))
            if text.strip():
                texts.append(text)
                indexes.append(self.p_index[element])
        # The lower cells of a vertical merge are empty stand-ins for the top one.
        attached = any(any(extras.values()) for extras, _ in extras_list)
        if continues and not texts and not nested and not attached:
            return None
        cell = self.doc.add(
            BlockType.TABLE_CELL,
            "\n".join(texts),
            container=Container.TABLE,
            parent_id=table.block_id,
            section_path=table.section_path,
            row=row,
            col=col,
            header_id=header_id,
            native_ref={"part": self.name, "p": indexes},
        )
        for extras, index in extras_list:
            self._emit_extras(extras, cell.block_id, {"part": self.name, "p": [index]})
        for tbl in nested:
            self._table(tbl, cell.block_id)
        return cell

    def _classify(self, p) -> tuple[int | None, bool]:
        """(heading level or None, is a list item), from the paragraph and its style."""
        props = p.find(_w("pPr"))
        style_id = None
        level = numbered = None
        if props is not None:
            style = props.find(_w("pStyle"))
            style_id = style.get(_VAL) if style is not None else None
            outline = props.find(_w("outlineLvl"))
            if outline is not None:
                level = _int(outline, 9) + 1
            num = props.find(_w("numPr"))
            if num is not None:
                numbered = _int(num.find(_w("numId")), 1) != 0  # numId 0 switches numbering off
        style_level, style_numbered = self.styles.get(style_id, (None, False))
        level = style_level if level is None else level
        if level is not None and level > 9:  # outline level 9 means "body text"
            level = None
        return level, style_numbered if numbered is None else numbered

    def _inline(self, p) -> tuple[str, dict]:
        """The paragraph's visible text, plus everything attached to it that is not."""
        out: list[str] = []
        extras = {"images": [], "boxes": [], "links": [], "fields": [], "deleted": [], "authors": []}

        def visit(element, deleted=False):
            for child in element:
                tag = child.tag
                if not isinstance(tag, str) or tag in _SKIP:
                    continue
                if tag == _w("t"):
                    (extras["deleted"] if deleted else out).append(child.text or "")
                elif tag == _w("delText"):
                    extras["deleted"].append(child.text or "")
                elif deleted:
                    visit(child, True)
                elif tag == _w("tab"):
                    out.append("\t")
                elif tag in (_w("br"), _w("cr")):
                    out.append("\n")
                elif tag == _w("noBreakHyphen"):
                    out.append("-")
                elif tag == _w("instrText"):
                    extras["fields"].append(child.text or "")
                elif tag in _GRAPHICS:
                    self._graphic(child, extras)
                elif tag == f"{{{_MC}}}AlternateContent":
                    # Two renderings of the same object for new and old Word; read one.
                    choice = child.find(f"{{{_MC}}}Choice")
                    visit(choice if choice is not None else child)
                else:
                    if tag in _DELETED or tag in _INSERTED:
                        extras["authors"].append(child.get(_AUTHOR))
                    if tag == _w("hyperlink") and child.get(f"{{{_R}}}id"):
                        rel = self.part.rels.get(child.get(f"{{{_R}}}id"))
                        if rel is not None and rel.is_external:
                            extras["links"].append(rel.target_ref)
                    visit(child, tag in _DELETED)

        visit(p)
        return "".join(out), extras

    def _graphic(self, element, extras: dict) -> None:
        alt = next((e.get("descr") for e in element.iter(f"{{{_WP}}}docPr") if e.get("descr")), None)
        for blip in element.iter(f"{{{_A}}}blip", f"{{{_V}}}imagedata"):
            rid = blip.get(f"{{{_R}}}embed") or blip.get(f"{{{_R}}}link") or blip.get(f"{{{_R}}}id")
            if rid:
                extras["images"].append((rid, alt))
        extras["boxes"].extend(element.iter(_w("txbxContent")))

    def _emit_extras(self, extras: dict, owner: str | None, ref: dict) -> None:
        def meta(text, flag):
            self.doc.add(
                BlockType.METADATA,
                text,
                parent_id=owner,
                section_path=self._section(),
                native_ref={**ref, flag: True},
            )

        for rid, alt in extras["images"]:
            image_ref = {**ref, "rid": rid}
            rel = self.part.rels.get(rid)
            target = None
            if rel is None or rel.is_external:  # linked, not embedded: no bytes to OCR
                image_ref["linked"] = True
            else:
                target = rel.target_part
                image_ref.update(
                    target=str(target.partname),
                    sha1=hashlib.sha1(target.blob).hexdigest(),
                    content_type=target.content_type,
                )
            image = self.doc.add(
                BlockType.IMAGE_REGION,
                container=Container.IMAGE,
                parent_id=owner,
                section_path=self._section(),
                native_ref=image_ref,
            )
            if alt and alt.strip():
                self.doc.add(
                    BlockType.METADATA,
                    alt,
                    parent_id=image.block_id,
                    section_path=self._section(),
                    native_ref={**ref, "alt_text": True},
                )
            if target is not None:
                reading = self.reader.read(target.blob, target.content_type)
                images.emit(self.doc, image, reading)
        for box in extras["boxes"]:
            self.walk(box, owner, BlockType.TEXT_BOX)
        for url in dict.fromkeys(extras["links"]):
            meta(url, "hyperlink")
        field = "".join(extras["fields"]).strip()
        if "HYPERLINK" in field:  # a link written as a field code; other fields carry no data
            meta(field, "hyperlink")
        if "".join(extras["deleted"]).strip():
            meta("".join(extras["deleted"]), "tracked_deletion")
        for author in dict.fromkeys(a for a in extras["authors"] if a):
            meta(author, "revision_author")

def _block_children(container):
    """Paragraphs and tables directly inside a container, seeing through wrappers
    (content controls, tracked insertions) that sit between them and it."""
    for child in container:
        if child.tag in (_P, _TBL):
            yield child
        elif isinstance(child.tag, str) and child.tag not in (_w("tcPr"), _w("sectPr"), _w("sdtPr")):
            if child.tag == _w("sdt"):
                content = child.find(_SDT_CONTENT)
                if content is not None:
                    yield from _block_children(content)
            elif child.tag in _INSERTED or child.tag in (_w("customXml"), _w("smartTag")):
                yield from _block_children(child)


def _style_table(styles_root) -> dict[str, tuple[int | None, bool]]:
    """style id -> (heading level, has numbering), following each style's parent."""
    raw = {}
    for style in styles_root.iter(_w("style")):
        props = style.find(_w("pPr"))
        outline = props.find(_w("outlineLvl")) if props is not None else None
        name = style.find(_w("name"))
        based = style.find(_w("basedOn"))
        raw[style.get(_w("styleId"))] = (
            (name.get(_VAL) or "").lower() if name is not None else "",
            _int(outline, 9) + 1 if outline is not None else None,
            props is not None and props.find(_w("numPr")) is not None,
            based.get(_VAL) if based is not None else None,
        )

    def resolve(style_id, seen=()):
        if style_id not in raw or style_id in seen:
            return None, False
        name, level, numbered, based = raw[style_id]
        parent_level, parent_numbered = resolve(based, (*seen, style_id))
        if level is None:
            level = 0 if name == "title" else parent_level
        return level, numbered or parent_numbered

    return {style_id: resolve(style_id) for style_id in raw}


def _int(element, default: int) -> int:
    try:
        return int(element.get(_VAL))
    except (AttributeError, TypeError, ValueError):
        return default


def _local(element) -> str:
    return element.tag.rsplit("}", 1)[-1]
