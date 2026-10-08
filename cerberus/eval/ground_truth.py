"""Ground truth: PII a person found by reading the document, loaded from a workbook.

The sheets in `pii.py` start from what the pipeline found, which is quick to
mark but biased: a person checking a list tends to miss what is not on it.
This is the independent answer. Someone reads the pages and types every PII
instance into a spreadsheet without looking at the pipeline's output, one row
per instance, one tab per person.

Columns are found by their heading, so their order and any extra columns
(a row number, a cross-check column) do not matter:

    File   Page   Section   Location   Type   Exact text   Notes

The workbook holds real values. It lives in `labels/` and is never committed.
"""

from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path

# Heading (lower case, as it starts) -> field. "Page / Slide" starts with "page".
_HEADINGS = {
    "file": "file", "page": "page", "section": "section", "location": "location",
    "type": "type", "exact text": "text", "notes": "notes",
}
_REQUIRED = ("page", "type", "text")


@dataclass(frozen=True)
class Label:
    row: int  # row in the sheet, from 1, so a result can be traced back to it
    page: int | None
    type: str  # the labeller's category, e.g. PERSON, EMAIL
    text: str  # exactly as the person read it
    location: str = ""  # body, table, prose or image
    section: str = ""
    file: str = ""
    notes: str = ""


def normalise(value: object) -> str:
    """Lower case, runs of white space as one space: how values are compared."""
    return " ".join(str(value).casefold().split())


def load(path: str | Path, sheet: str) -> list[Label]:
    """Every labelled instance on one tab of the workbook, in sheet order."""
    import openpyxl  # here, so the rest of the package does not need it

    workbook = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
    try:
        if sheet not in workbook.sheetnames:
            raise ValueError(f"no tab named {sheet!r}; the workbook has {', '.join(workbook.sheetnames)}")
        rows = list(workbook[sheet].iter_rows(values_only=True))
    finally:
        workbook.close()

    columns: dict[str, int] = {}
    first = 0
    for first, row in enumerate(rows[:10]):
        columns = _columns(row)
        if all(name in columns for name in _REQUIRED):
            break
    else:
        raise ValueError(f"tab {sheet!r} has no heading row with Page, Type and Exact text")

    def cell(row: tuple, name: str) -> str:
        index = columns.get(name)
        value = row[index] if index is not None and index < len(row) else None
        return "" if value is None else str(value).strip()

    labels = []
    for number, row in enumerate(rows[first + 1 :], start=first + 2):
        text = cell(row, "text")
        if not text:
            continue  # an empty or unfinished row
        page = cell(row, "page")
        labels.append(Label(
            row=number, page=int(float(page)) if page.replace(".", "", 1).isdigit() else None,
            type=cell(row, "type").upper(), text=text, location=cell(row, "location").lower(),
            section=cell(row, "section"), file=cell(row, "file"), notes=cell(row, "notes"),
        ))
    return labels


def _columns(row: tuple) -> dict[str, int]:
    found: dict[str, int] = {}
    for index, value in enumerate(row):
        heading = normalise(value or "")
        for start, name in _HEADINGS.items():
            if heading.startswith(start):
                found.setdefault(name, index)
    return found
