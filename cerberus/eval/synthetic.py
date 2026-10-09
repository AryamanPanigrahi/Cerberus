"""Small made-up documents that reproduce the difficult cases by format, not by value.

Every name, number and address here is invented. The files are safe to
commit, to show on a slide and to run in front of an audience, and they are
small enough that the whole pipeline runs on them in about a minute.

What each one exercises:

    staff_notice.docx   typed text, a table with a person column, a staff number
                        in running prose, a picture of an access card (name,
                        number, address and a barcode as pixels only), an
                        author in the file properties
    incident_scan.pdf   a scanned page: no text layer, so everything is OCR;
                        identifiers in unlabelled prose, a card number, a
                        tax number, a test-range social security number, an
                        author stored in the PDF's properties
    team.pptx           separate text boxes read in the wrong order by a naive
                        reader, and a table that uses abbreviated names

Run:  python -m cerberus.eval.synthetic tests/fixtures/synthetic
"""

from __future__ import annotations

import io
from pathlib import Path

PEOPLE = [  # name, staff number, e-mail, mobile
    ("Jane Sample", "EMP-00001", "j.sample@corp.example", "+1 (555) 010-0101"),
    ("Omar Placeholder", "EMP-00002", "o.placeholder@corp.example", "+1 (555) 010-0102"),
    ("Ana Example", "EMP-00003", "a.example@corp.example", "+44 20 7946 0103"),
]
AUTHOR = "Maria Fixture"  # stored only in each file's hidden properties
CARD = "4111 1111 1111 1111"  # a published test number: passes the Luhn check
SSN = "900-12-3456"  # the 900 range is never issued
PAN = "ABCPS1234D"  # right structure; fourth letter P marks an individual


def make(folder: str | Path) -> list[Path]:
    """Write the three sample files into folder and return their paths."""
    folder = Path(folder)
    folder.mkdir(parents=True, exist_ok=True)
    return [_docx(folder / "staff_notice.docx"), _scanned_pdf(folder / "incident_scan.pdf"), _pptx(folder / "team.pptx")]


def _font(size: int):
    from PIL import ImageFont

    return ImageFont.load_default(size=size)


def access_card() -> bytes:
    """A picture of an access card: its text and barcode exist only as pixels."""
    from PIL import Image, ImageDraw

    name, number, mail, _ = PEOPLE[1]
    card = Image.new("RGB", (900, 420), "white")
    draw = ImageDraw.Draw(card)
    draw.rectangle((0, 0, 900, 90), fill=(24, 44, 92))
    draw.text((30, 24), "SAMPLE CORPORATION", fill="white", font=_font(40))
    for row, (label, value) in enumerate((("NAME", name), ("STAFF NUMBER", number), ("WORK E-MAIL", mail))):
        top = 115 + row * 80
        draw.text((40, top), label, fill=(110, 110, 110), font=_font(18))
        draw.text((40, top + 24), value, fill="black", font=_font(34))
    x, widths = 560, (3, 5, 2, 6, 3, 4)
    while x < 850:  # a barcode nobody can decode, which is how most printed ones scan
        bar = widths[x % len(widths)]
        draw.rectangle((x, 300, x + bar, 385), fill="black")
        x += bar + widths[(x + 2) % len(widths)]
    out = io.BytesIO()
    card.save(out, format="PNG")
    return out.getvalue()


def _docx(path: Path) -> Path:
    import docx
    from docx.shared import Inches

    d = docx.Document()
    d.core_properties.author = AUTHOR
    d.core_properties.last_modified_by = AUTHOR
    d.add_heading("1. Staff notice", level=1)
    name, number, mail, phone = PEOPLE[0]
    d.add_paragraph(
        f"The new access procedure was approved by {name} ({number}) on Monday. Questions go to "
        f"{mail} or {phone}. The primary supplier contact and the e-mail of the supplier lead are "
        "mandatory fields on the request form."
    )
    d.add_heading("1.1 Named approvers", level=2)
    table = d.add_table(rows=1, cols=4)
    table.style = "Table Grid"
    for cell, text in zip(table.rows[0].cells, ("Approver", "Staff ID", "E-mail", "Mobile")):
        cell.text = text
    for person in PEOPLE:
        for cell, text in zip(table.add_row().cells, person):
            cell.text = text
    d.add_heading("1.2 Sample access card", level=2)
    d.add_paragraph("Figure 1 shows a card as issued.")
    d.add_picture(io.BytesIO(access_card()), width=Inches(4.5))
    d.save(str(path))
    return path


def _scanned_pdf(path: Path) -> Path:
    """One page that is a single image, as a scanner would produce it."""
    import pymupdf
    from PIL import Image, ImageDraw

    name, number, _, phone = PEOPLE[2]
    lines = [
        ("4. Incident summary", 44),
        ("", 20),
        ("On the morning of the review the duty manager spoke to", 30),
        (f"{name}, who confirmed that the card ending in the", 30),
        (f"number {CARD} had been used twice. Her staff", 30),
        (f"number is {number} and she can be reached on", 30),
        (f"{phone}. The claimant gave the tax number {PAN}", 30),
        (f"and a social security number of {SSN} on the form.", 30),
        ("", 20),
        ("No further action was taken by the service desk that day.", 30),
    ]
    page = Image.new("L", (1654, 2339), 255)  # A4 at 200 dpi
    draw = ImageDraw.Draw(page)
    top = 220
    for text, size in lines:
        draw.text((180, top), text, fill=0, font=_font(size))
        top += int(size * 1.9)
    scan = io.BytesIO()
    page.save(scan, format="PNG")

    pdf = pymupdf.open()
    sheet = pdf.new_page(width=595, height=842)
    sheet.insert_image(sheet.rect, stream=scan.getvalue())
    pdf.set_metadata({"author": AUTHOR, "title": "Incident summary"})
    pdf.save(str(path))
    pdf.close()
    return path


def _pptx(path: Path) -> Path:
    from pptx import Presentation
    from pptx.util import Emu, Pt

    prs = Presentation()
    prs.core_properties.author = AUTHOR
    prs.core_properties.last_modified_by = AUTHOR  # the library's template carries its own author's name here
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    slide.shapes.title.text = "Programme team"
    # Added right to left: a reader that takes shapes in stored order gets them backwards.
    for column, (name, _, mail, phone) in reversed(list(enumerate(PEOPLE))):
        box = slide.shapes.add_textbox(Emu(500_000 + column * 2_900_000), Emu(1_700_000), Emu(2_700_000), Emu(1_100_000))
        frame = box.text_frame
        frame.text = name
        frame.paragraphs[0].runs[0].font.bold = True
        for extra in (mail, phone):
            frame.add_paragraph().text = extra
        for paragraph in frame.paragraphs:
            for run in paragraph.runs:
                run.font.size = Pt(14)
    rows = [("Activity", "Responsible", "Accountable"), ("Access review", "J. Sample", "O. Placeholder"),
            ("Supplier checks", "A. Example", "J. Sample")]
    shape = slide.shapes.add_table(len(rows), 3, Emu(500_000), Emu(3_300_000), Emu(8_000_000), Emu(1_400_000))
    for r, row in enumerate(rows):
        for c, text in enumerate(row):
            shape.table.cell(r, c).text = text
    prs.save(str(path))
    return path


if __name__ == "__main__":
    import sys

    for written in make(sys.argv[1] if len(sys.argv) > 1 else "tests/fixtures/synthetic"):
        print("wrote", written)
