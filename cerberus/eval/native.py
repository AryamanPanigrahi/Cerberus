"""Structure retention for DOCX and PPTX, over the whole file, with no hand-made reference.

A scan has no record of its own structure, so a person has to write one down
(structure.py). A DOCX or PPTX does have one: the file says which paragraph is
a heading, which cell is in which row and column, which shape is on which
slide. This module reads that record straight from the file, with a few lines
of python-docx / python-pptx that share no code with the extractors, and
checks the extraction against it. Every element of the file is counted, so
the figure is "N of M elements, whole file", not a sample.

What counts as retained (the definition in structure.py):

    the element is in the extraction, with its text, as the right kind of
    element, under the right parent, and in the right order.

- DOCX. The file fixes the order, so order is scored: elements are paired by
  the longest run of matching texts that keeps both sequences in order. The
  parent of a paragraph is the heading above it; of a cell, its column header
  and the first cell of its row.
- PPTX. A slide is shapes placed on a canvas; the file stores them in drawing
  order, which is not reading order, and nothing in it says what the reading
  order is. So order is not scored for shapes: an element is retained if it
  is on the right slide as the right kind. Table cells are scored as in DOCX.

Left out of the reference: text boxes, headers, footers and footnotes in a
DOCX, and tables nested inside a cell. They are extracted, but the few lines
here do not enumerate them.

Run:  python -m cerberus.eval.native FILE [FILE ...]
"""

from __future__ import annotations

import argparse
import re
from pathlib import Path

from . import structure
from .structure import Element, Score

_W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"


# --- DOCX ---------------------------------------------------------------------


def docx_reference(path: str | Path) -> list[Element]:
    """Every body paragraph and table cell of a DOCX, in file order, with its kind and parent.

    Three things in the format that a first, simpler version of this got wrong
    (it then blamed the extractor for 40 elements that were in fact right):

    - a heading is a paragraph with an outline level, set on the paragraph or
      inherited from its style. Looking only for styles named "Heading N"
      misses a paragraph given an outline level directly;
    - a cell's text can sit inside a content control, a field or an inserted
      revision, none of which python-docx's `.text` reads;
    - a cell merged down several rows is stored once, with empty placeholders
      below it. For the rows below, it is still the first cell of the row.
    """
    import docx
    from docx.text.paragraph import Paragraph

    source = docx.Document(str(path))
    items: list[Element] = []
    levels: list[int] = []  # heading levels currently open, outermost first
    tables = 0

    def text_of(cell) -> str:
        own = [p for p in cell.iter(f"{_W}p") if next(a for a in p.iterancestors() if a.tag == f"{_W}tc") is cell]
        return structure._one_line(" ".join(_text(p) for p in own))

    def walk(parent) -> None:
        nonlocal tables
        for child in parent.iterchildren():
            if child.tag == f"{_W}sdt":  # a content control: its contents are ordinary body content
                content = child.find(f"{_W}sdtContent")
                if content is not None:
                    walk(content)
            elif child.tag == f"{_W}p":
                paragraph = Paragraph(child, source)
                text = structure._one_line(_text(child))
                level = _outline_level(child, paragraph)
                if level is not None and text:
                    while levels and levels[-1] >= level:
                        levels.pop()
                    levels.append(level)
                    items.append(Element(f"h{len(levels)}", text))
                elif _numbered(child, paragraph):
                    items.append(Element("list_item", text))
                else:
                    items.append(Element("paragraph", text))
            elif child.tag == f"{_W}tbl":
                above: dict[int, str] = {}  # column -> text of the merged cell that reaches down into it
                for r, row in enumerate(child.iterchildren(f"{_W}tr")):
                    column = 0
                    for cell in row.iterchildren(f"{_W}tc"):
                        properties = cell.find(f"{_W}tcPr")
                        merge = properties.find(f"{_W}vMerge") if properties is not None else None
                        span = properties.find(f"{_W}gridSpan") if properties is not None else None
                        width = int(span.get(f"{_W}val")) if span is not None else 1
                        text = text_of(cell)
                        if merge is not None and merge.get(f"{_W}val") != "restart" and not text:
                            item = Element("cell", above.get(column, ""), stand_in=True, table=tables, row=r, col=column)
                        else:
                            item = Element("cell", text, table=tables, row=r, col=column)
                            if merge is not None:
                                above[column] = text
                            else:
                                above.pop(column, None)
                        items.append(item)
                        column += width
                tables += 1

    walk(source.element.body)
    return structure._link(items)


def _text(paragraph) -> str:
    """The paragraph's own visible text, read from the XML.

    python-docx's `.text` sees plain runs and hyperlinks and nothing else, so
    a cell whose text sits in a field or an inserted revision reads as empty.
    Here every text node counts, except deleted text (a different tag) and
    what belongs to a text box drawn inside the paragraph, which is a
    separate element of the document.
    """
    skipped = (f"{_W}txbxContent", "{http://schemas.openxmlformats.org/markup-compatibility/2006}Fallback",
               f"{_W}pict", f"{_W}instrText")
    parts = []
    for node in paragraph.iter(f"{_W}t", f"{_W}tab", f"{_W}br", f"{_W}cr"):
        inside = False
        for ancestor in node.iterancestors():
            if ancestor is paragraph:
                break
            if ancestor.tag in skipped:
                inside = True
                break
        if not inside:
            parts.append(node.text or "" if node.tag == f"{_W}t" else " ")
    return "".join(parts)


def _outline_level(element, paragraph) -> int | None:
    """The paragraph's heading level, from its own outline level or its style's; None for body text."""
    def read(properties):
        outline = properties.find(f"{_W}outlineLvl") if properties is not None else None
        return int(outline.get(f"{_W}val")) if outline is not None else None

    level = read(element.find(f"{_W}pPr"))
    style = paragraph.style
    while level is None and style is not None:
        level = read(style.element.find(f"{_W}pPr"))
        if level is None and (style.name or "").lower() == "title":
            level = -1  # a title outranks every heading
        style = style.base_style
    return None if level is None or level >= 9 else level + 1  # 9 is the format's "body text"


def _numbered(element, paragraph) -> bool:
    """Is the paragraph a list item: numbered or bulleted directly, or through its style?"""
    if element.find(f"{_W}pPr/{_W}numPr") is not None:
        return True
    style = paragraph.style
    while style is not None:
        properties = style.element.find(f"{_W}pPr")
        if properties is not None and properties.find(f"{_W}numPr") is not None:
            return True
        style = style.base_style
    return False


# --- PPTX ---------------------------------------------------------------------


def pptx_reference(path: str | Path) -> dict[int, list[Element]]:
    """For each slide: its text-bearing shapes and table cells, with kind and parent."""
    from pptx import Presentation
    from pptx.shapes.group import GroupShape

    slides: dict[int, list[Element]] = {}
    for number, slide in enumerate(Presentation(str(path)).slides, start=1):
        title = slide.shapes.title
        items: list[Element] = []
        tables = 0

        def walk(shapes) -> None:
            nonlocal tables
            for shape in shapes:
                if isinstance(shape, GroupShape):
                    walk(shape.shapes)
                elif getattr(shape, "has_table", False) and shape.has_table:
                    for r, row in enumerate(shape.table.rows):
                        for c, cell in enumerate(row.cells):
                            items.append(Element("cell", structure._one_line(cell.text_frame.text), table=tables, row=r, col=c))
                    tables += 1
                elif shape.has_text_frame and shape.text_frame.text.strip():
                    is_title = title is not None and shape.shape_id == title.shape_id
                    items.append(Element("h1" if is_title else "paragraph", structure._one_line(shape.text_frame.text)))

        walk(slide.shapes)
        linked = structure._link(items)  # gives cells their header and row
        for item in linked:
            if item.kind != "cell":
                item.parent = ("slide", str(number))
        slides[number] = linked
    return slides


# --- scoring ------------------------------------------------------------------


def score_docx(path: str | Path, doc) -> Score:
    """Whole-file score of a DOCX extraction against the file's own structure. Order is scored."""
    reference = docx_reference(path)
    extracted = [e for e in structure.elements(doc, doc.blocks) if not e.image]
    return structure.score(reference, extracted)


def score_pptx(path: str | Path, doc) -> Score:
    """Whole-file score of a PPTX extraction. Shapes: right slide and kind. Cells: right header and row."""
    total = Score()
    by_slide = structure.units(doc)
    for number, reference in pptx_reference(path).items():
        extracted = [e for e in structure.elements(doc, by_slide.get(number, [])) if not e.image]
        for item in extracted:
            if item.kind != "cell":
                item.parent = ("slide", str(number))
        total.add(_unordered(reference, extracted))
    return total


def _unordered(reference: list[Element], extracted: list[Element]) -> Score:
    """Pair elements by text alone, best match first, each extracted element used once."""
    result = Score(total=len(reference))
    free = list(range(len(extracted)))
    for ref in reference:
        kind = "heading" if ref.kind.startswith("h") else ref.kind
        tally = result.by_kind.setdefault(kind, [0, 0])
        tally[1] += 1
        scored = [(structure._similarity(structure._norm(ref.text), structure._norm(extracted[i].text)), -i) for i in free]
        # Among equally good texts (the same name in three cells) prefer one that is also in the right place.
        scored = [(s, extracted[-i].kind == ref.kind and structure._same_parent(ref.parent, extracted[-i].parent), i)
                  for s, i in scored if s >= structure.SIMILAR]
        if not scored:
            result.missing += 1
            continue
        similarity, _, negative = max(scored)
        got = extracted[-negative]
        free.remove(-negative)
        result.characters += similarity * len(ref.text)
        result.length += len(ref.text)
        if got.kind != ref.kind:
            result.wrong_kind += 1
        elif not structure._same_parent(ref.parent, got.parent):
            result.wrong_parent += 1
        else:
            result.retained += 1
            tally[0] += 1
    return result


def report(name: str, result: Score, ordered: bool) -> str:
    lines = [
        f"{name}: {result.retained} of {result.total} elements retained = {result.share:.1%} "
        f"(whole file; order {'scored' if ordered else 'not scored: a slide has no reading order of its own'})",
        f"  lost: {result.missing} missing or out of order, {result.wrong_kind} wrong kind, {result.wrong_parent} wrong parent",
        f"  character accuracy of matched elements: {result.character_accuracy:.1%}",
    ]
    lines += [f"  {kind:<10} {kept} of {count}" for kind, (kept, count) in sorted(result.by_kind.items())]
    return "\n".join(lines)


def main(argv: list[str] | None = None) -> None:
    from ..extract import extract

    parser = argparse.ArgumentParser(description=__doc__.split("\n")[0])
    parser.add_argument("files", nargs="+", help="DOCX and PPTX files")
    args = parser.parse_args(argv)
    for path in map(Path, args.files):
        doc = extract(path)
        if doc.format == "docx":
            print(report(path.name, score_docx(path, doc), ordered=True))
        elif doc.format == "pptx":
            print(report(path.name, score_pptx(path, doc), ordered=False))
        else:
            print(f"{path.name}: a {doc.format} has no structure of its own to read; use cerberus.eval.structure")


if __name__ == "__main__":
    main()
