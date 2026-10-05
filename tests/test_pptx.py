import io

import pytest
from PIL import Image
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.util import Emu, Pt

from cerberus.extract.ir import BlockType, Container
from cerberus.extract.pptx import extract_pptx

# All names and contact details below are invented.


def box(slide, left, top, width, height, text="", size=None):
    shape = slide.shapes.add_textbox(Emu(left), Emu(top), Emu(width), Emu(height))
    if text:
        lines = text.split("\n")
        shape.text_frame.text = lines[0]
        for line in lines[1:]:
            shape.text_frame.add_paragraph().text = line
        for paragraph in shape.text_frame.paragraphs:
            for run in paragraph.runs:
                run.font.size = Pt(size or 10)
    return shape


def card(slide, left, top, name, role, contact, scrambled=False):
    """Background rectangle plus three text boxes, like a hand-drawn org-chart card."""
    parts = [
        lambda: box(slide, left, top, 2_000_000, 1_000_000),  # background, no text
        lambda: box(slide, left + 100_000, top + 50_000, 1_800_000, 200_000, name),
        lambda: box(slide, left + 100_000, top + 250_000, 1_800_000, 200_000, role),
        lambda: box(slide, left + 100_000, top + 450_000, 1_800_000, 400_000, contact),
    ]
    for make in reversed(parts) if scrambled else parts:
        make()


def texts(blocks):
    return [b.text for b in blocks]


@pytest.fixture
def deck(tmp_path):
    prs = Presentation()
    prs.core_properties.author = "Jane Sample"
    blank = prs.slide_layouts[6]

    # Slide 1: title drawn as a plain box, two cards added right-to-left and
    # inside-out so drawing order disagrees with reading order.
    s1 = prs.slides.add_slide(blank)
    card(s1, 5_000_000, 2_000_000, "Omar Placeholder", "HEAD OF TESTING",
         "o.placeholder@corp.example\nM +1 (555) 0100", scrambled=True)
    card(s1, 1_000_000, 2_000_000, "Jane Sample", "HEAD OF EXAMPLES",
         "j.sample@corp.example\nM +1 (555) 0101")
    box(s1, 500_000, 200_000, 6_000_000, 500_000, "Team Structure", size=28)
    s1.notes_slide.notes_text_frame.text = "Ask Jane Sample about the rollout."

    # Slide 2: a table, a picture, and a group shape that is later stretched.
    s2 = prs.slides.add_slide(blank)
    frame = s2.shapes.add_table(3, 2, Emu(500_000), Emu(500_000), Emu(4_000_000), Emu(900_000))
    for r, row in enumerate([("Name", "Mobile"), ("Jane Sample", "+1 (555) 0101"),
                             ("Omar Placeholder", "+1 (555) 0100")]):
        for c, value in enumerate(row):
            frame.table.cell(r, c).text = value
    png = io.BytesIO()
    Image.new("RGB", (40, 20), "white").save(png, format="PNG")
    png.seek(0)
    s2.shapes.add_picture(png, Emu(6_000_000), Emu(500_000), Emu(800_000), Emu(400_000))
    group = s2.shapes.add_group_shape()
    inner = group.shapes.add_textbox(Emu(1_000_000), Emu(3_000_000), Emu(1_000_000), Emu(300_000))
    inner.text_frame.text = "Grouped text"
    group.width = group.width * 2  # children keep their stored coordinates

    # Slide 3: a chart, which the extractor cannot read.
    s3 = prs.slides.add_slide(blank)
    data = CategoryChartData()
    data.categories = ["a", "b"]
    data.add_series("s", (1, 2))
    s3.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Emu(500_000), Emu(500_000),
                        Emu(3_000_000), Emu(2_000_000), data)

    path = tmp_path / "deck.pptx"
    prs.save(str(path))
    return extract_pptx(path)


def on_slide(doc, number, type_=None):
    return [b for b in doc.blocks if b.page == number and (type_ is None or b.type is type_)]


def test_metadata_and_pages(deck):
    meta = [b for b in deck.blocks if b.type is BlockType.METADATA and b.page is None]
    # python-pptx's default template fills in a few properties of its own.
    assert ("Jane Sample", "author") in [(b.text, b.native_ref["core_property"]) for b in meta]
    assert [p.number for p in deck.pages] == [1, 2, 3]
    assert deck.validate() == []


def test_title_is_largest_text_and_sets_section_path(deck):
    (title,) = on_slide(deck, 1, BlockType.HEADING)
    assert title.text == "Team Structure"
    others = [b for b in on_slide(deck, 1) if b.type not in (BlockType.SLIDE, BlockType.HEADING)]
    assert all(b.section_path == ("Team Structure",) for b in others)


def test_cards_are_grouped_and_read_left_to_right(deck):
    groups = on_slide(deck, 1, BlockType.GROUP)
    assert len(groups) == 2
    assert texts(deck.children(groups[0].block_id)) == [
        "Jane Sample", "HEAD OF EXAMPLES", "j.sample@corp.example\nM +1 (555) 0101",
    ]
    assert texts(deck.children(groups[1].block_id))[0] == "Omar Placeholder"
    # Title first, then the left card, then the right card.
    order = [b.text for b in on_slide(deck, 1) if b.text]
    assert order.index("Team Structure") < order.index("Jane Sample") < order.index("Omar Placeholder")


def test_notes_are_extracted(deck):
    (notes,) = on_slide(deck, 1, BlockType.NOTES)
    assert notes.text == "Ask Jane Sample about the rollout."


def test_table_cells_point_at_their_header(deck):
    (table,) = on_slide(deck, 2, BlockType.TABLE)
    cells = deck.children(table.block_id)
    assert texts(cells) == ["Name", "Mobile", "Jane Sample", "+1 (555) 0101",
                            "Omar Placeholder", "+1 (555) 0100"]
    assert all(c.container is Container.TABLE for c in cells)
    assert [deck.get(c.header_id).text if c.header_id else None for c in cells] == [
        None, None, "Name", "Mobile", "Name", "Mobile",
    ]
    assert (cells[3].row, cells[3].col) == (1, 1)
    assert cells[3].bbox.x0 == pytest.approx(2_500_000) and cells[3].bbox.y0 > cells[0].bbox.y0


def test_picture_becomes_image_region(deck):
    (image,) = on_slide(deck, 2, BlockType.IMAGE_REGION)
    assert image.container is Container.IMAGE
    assert image.native_ref["content_type"] == "image/png"
    assert image.bbox.x0 == 6_000_000


def test_group_shape_children_use_slide_coordinates(deck):
    (group,) = on_slide(deck, 2, BlockType.GROUP)
    (inner,) = deck.children(group.block_id)
    assert inner.text == "Grouped text"
    # The group was stretched to double width, so the child is too.
    assert inner.bbox.x0 == pytest.approx(1_000_000)
    assert inner.bbox.x1 - inner.bbox.x0 == pytest.approx(2_000_000)


def test_unreadable_content_quarantines_the_slide(deck):
    assert [p.quarantined for p in deck.pages] == [False, False, True]
    assert "chart" in deck.pages[2].reason
