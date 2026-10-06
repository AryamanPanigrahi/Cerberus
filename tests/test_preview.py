import docx

from cerberus.report import preview

# All names and contact details below are invented.


def test_preview_highlights_findings_and_escapes_document_text(tmp_path):
    d = docx.Document()
    d.add_heading("1. Contacts", level=1)
    d.add_paragraph("Reach <Jane> at j.sample@corp.example or +1 (555) 010-0101.")
    table = d.add_table(rows=2, cols=2)
    for r, row in enumerate([("NAME", "EMP ID"), ("Jane Sample", "EMP-00001")]):
        for c, value in enumerate(row):
            table.cell(r, c).text = value
    source = tmp_path / "sample.docx"
    d.save(str(source))

    index, page = preview.build([source], tmp_path / "out")
    html = page.read_text(encoding="utf-8")

    assert "sample.html" in index.read_text(encoding="utf-8")
    assert "j.sample@corp.example<sub>EMAIL</sub></mark>" in html
    assert "<sub>PHONE</sub>" in html and "EMP-00001<sub>STAFF_ID</sub>" in html
    assert "&lt;Jane&gt;" in html and "<Jane>" not in html  # document text cannot inject markup
    assert "<th>NAME</th>" in html  # the table is drawn as a table, header row first
    assert "Findings by entity and location" in html
