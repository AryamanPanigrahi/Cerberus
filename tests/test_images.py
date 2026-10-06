import io
import os
import shutil

import docx
import pymupdf
import pytest
from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.util import Emu

from cerberus.extract import images, ocr
from cerberus.extract.docx import extract_docx
from cerberus.extract.ir import BlockType, Container, Document, Source, Unit
from cerberus.extract.pdf import extract_pdf
from cerberus.extract.pptx import extract_pptx

# All names and identifiers below are invented.

try:
    ocr.engine_version()
    HAVE_TESSERACT = True
except RuntimeError:
    HAVE_TESSERACT = False
needs_tesseract = pytest.mark.skipif(not HAVE_TESSERACT, reason="Tesseract is not installed")
HAVE_SOFFICE = bool(
    os.environ.get("SOFFICE_CMD") or shutil.which("soffice") or images._WINDOWS_SOFFICE.exists()
)


def screenshot() -> bytes:
    """A small picture of a user record, as a PNG."""
    picture = Image.new("RGB", (640, 200), "white")
    draw = ImageDraw.Draw(picture)
    font = ImageFont.load_default(size=28)
    draw.text((30, 30), "User: Jane Sample", fill="black", font=font)
    draw.text((30, 110), "Badge VND-00002", fill="black", font=font)
    out = io.BytesIO()
    picture.save(out, format="PNG")
    return out.getvalue()


def image_text(doc: Document) -> list[str]:
    return [b.text for b in doc.blocks if b.container is Container.IMAGE and b.type is BlockType.PARAGRAPH]


def test_a_picture_that_cannot_be_opened_is_quarantined():
    doc = Document("d", "x", "docx")
    region = doc.add(BlockType.IMAGE_REGION, container=Container.IMAGE)
    images.emit(doc, region, images.Reader().read(b"not a picture", "image/png"))
    assert region.quarantined and "could not open" in region.native_ref["unreadable"]
    assert len(doc.blocks) == 1


def test_an_icon_sized_picture_is_not_read():
    out = io.BytesIO()
    Image.new("RGB", (24, 24), "white").save(out, format="PNG")
    reading = images.Reader().read(out.getvalue(), "image/png")
    assert reading.lines == [] and reading.conf is None and not reading.error


@needs_tesseract
def test_text_in_a_docx_picture_becomes_image_content(tmp_path):
    d = docx.Document()
    d.add_heading("1. Evidence", level=1)
    d.add_picture(io.BytesIO(screenshot()))
    path = tmp_path / "doc.docx"
    d.save(str(path))
    doc = extract_docx(path)

    (region,) = [b for b in doc.blocks if b.type is BlockType.IMAGE_REGION]
    lines = doc.children(region.block_id)
    assert [b.text for b in lines] == ["User: Jane Sample", "Badge VND-00002"]
    assert region.source is Source.OCR and region.ocr_conf > 0.8 and not region.quarantined
    assert all(b.section_path == ("1. Evidence",) and b.source is Source.OCR for b in lines)
    # Word boxes are in the picture's own pixels: where the text was drawn.
    box = lines[1].span_boxes(lines[1].text.index("VND"), len(lines[1].text))[0]
    assert box.unit is Unit.PX and 100 < box.x0 < 160 and 100 < box.y0 < 125
    assert doc.validate() == []


@needs_tesseract
def test_a_picture_used_twice_is_read_once(tmp_path, monkeypatch):
    calls = []
    real = images._read
    monkeypatch.setattr(images, "_read", lambda *a: calls.append(1) or real(*a))

    prs = Presentation()
    for _ in range(2):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        slide.shapes.add_picture(io.BytesIO(screenshot()), Emu(500_000), Emu(500_000))
    path = tmp_path / "deck.pptx"
    prs.save(str(path))
    doc = extract_pptx(path)

    assert len(calls) == 1
    assert image_text(doc) == ["User: Jane Sample", "Badge VND-00002"] * 2
    assert [b.page for b in doc.blocks if b.container is Container.IMAGE and b.text] == [1, 1, 2, 2]


@needs_tesseract
def test_picture_on_a_typed_pdf_page_is_read_and_placed_on_the_page(tmp_path):
    pdf = pymupdf.open()
    page = pdf.new_page()
    page.insert_text((72, 90), "The access record is reproduced below for reference only.", fontsize=11)
    page.insert_image(pymupdf.Rect(100, 200, 420, 300), stream=screenshot())  # half size
    path = tmp_path / "typed.pdf"
    pdf.save(str(path))
    doc = extract_pdf(path)

    assert doc.ocr_pages() == []  # the page itself has a text layer
    (region,) = [b for b in doc.blocks if b.type is BlockType.IMAGE_REGION]
    lines = doc.children(region.block_id)
    assert [b.text for b in lines] == ["User: Jane Sample", "Badge VND-00002"]
    box = lines[1].span_boxes(lines[1].text.index("VND"), len(lines[1].text))[0]
    assert box.unit is Unit.PT and 150 < box.x0 < 180 and 250 < box.y0 < 265


@needs_tesseract
@pytest.mark.skipif(not HAVE_SOFFICE, reason="LibreOffice is not installed")
def test_a_vector_picture_is_converted_before_reading():
    svg = (
        '<svg xmlns="http://www.w3.org/2000/svg" width="640" height="160">'
        '<rect width="640" height="160" fill="white"/>'
        '<text x="30" y="90" font-family="Arial" font-size="40">Jane Sample</text></svg>'
    )
    reading = images.Reader().read(svg.encode(), "image/svg+xml")
    assert not reading.error
    assert "Jane Sample" in " ".join(line.text for line in reading.lines)
