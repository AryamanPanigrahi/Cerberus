import io

import pytest
from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.util import Emu

from cerberus import pipeline
from cerberus.detect import merge
from cerberus.detect.finding import make
from cerberus.extract import images, neural, regions
from cerberus.extract.ir import BBox, BlockType, Container, Document, Source, Unit, Word
from cerberus.extract.ocr import OcrWord
from cerberus.extract.pptx import extract_pptx
from cerberus.redact import office, tokens, verify
from cerberus.redact.pptx import mask_pptx

# All names and identifiers below are invented.


def word(text, x0, y0, x1, y1, conf=0.9):
    return OcrWord(text, x0, y0, x1, y1, conf, (0, 0, 0))


# --- the second engine: tiles ---------------------------------------------


def test_tiles_cover_the_picture_and_overlap():
    assert neural.tiles(800, 600) == [(0, 0, 800, 600)]
    boxes = neural.tiles(2000, 900, size=1024, overlap=384)
    assert [b[0] for b in boxes] == [0, 640, 976] and all(b[2] - b[0] == 1024 for b in boxes)
    assert boxes[-1][2] == 2000 and all((b[1], b[3]) == (0, 900) for b in boxes)


def test_a_word_cut_at_a_tile_edge_gives_way_to_the_whole_word_from_the_next_tile():
    left, right = (0, 0, 1024, 500), (640, 0, 1664, 500)
    found = [
        (word("j.sample@corp.exa", 800, 100, 1024, 130), left),  # runs into the tile's right edge
        (word("j.sample@corp.example", 800, 100, 1080, 130), right),
        (word("Badge", 100, 100, 180, 130), left),
        (word("EMP-00001", 700, 300, 860, 330), left), (word("EMP-00001", 701, 301, 861, 331), right),
    ]
    merged = neural.merge(found, (1664, 500))
    assert sorted(w.text for w in merged) == ["Badge", "EMP-00001", "j.sample@corp.example"]
    # Words on one baseline share a line number; the next row down gets another.
    lines = {w.text: w.line for w in merged}
    assert lines["Badge"] == lines["j.sample@corp.example"] != lines["EMP-00001"]


# --- two readings compared -------------------------------------------------


def test_words_both_engines_read_alike_are_confirmed_whatever_the_word_breaks():
    first = [word("Jane", 10, 10, 60, 30), word("Sample,", 70, 10, 150, 30), word("j.sample@corp.example", 10, 50, 300, 70)]
    second = [word("Jane", 11, 11, 61, 31), word("Sample", 71, 11, 148, 31),
              word("j.sample", 10, 50, 110, 70), word("@corp.example", 112, 50, 300, 70)]
    agreement, doubtful = images.compare(first, second)
    assert agreement == 1.0 and doubtful == []


def test_a_word_only_one_engine_read_or_read_differently_is_doubtful():
    first = [word("Jane", 10, 10, 60, 30), word("Sampie", 70, 10, 150, 30)]
    second = [word("Jane", 11, 11, 61, 31), word("Sample", 71, 11, 148, 31), word("Placeholder", 10, 50, 160, 70)]
    agreement, doubtful = images.compare(first, second)
    assert sorted(w.text for w in doubtful) == ["Placeholder", "Sampie", "Sample"]
    assert agreement == pytest.approx(8 / 31)  # only "Jane", twice, out of every letter read
    assert images.compare([], []) == (None, [])


def test_a_picture_is_trusted_only_when_two_engines_agree_on_enough_of_it():
    lines = [object()]
    assert not images.Reading(lines, 0.9).trusted  # one engine cannot check itself
    assert images.Reading(lines, 0.9, second=[], agreement=0.8).trusted
    assert not images.Reading(lines, 0.9, second=[], agreement=0.3).trusted
    assert not images.Reading(lines, 0.4, second=[], agreement=0.9).trusted
    assert images.Reading([], None, second=[]).trusted  # both found no text at all
    assert not images.Reading(error="could not open picture", second=[]).trusted


def test_a_reading_is_reused_from_the_cache(tmp_path, monkeypatch):
    calls = []
    monkeypatch.setattr(images, "_examine", lambda gray, border, colour: calls.append(1) or images.Reading(size=gray.size))
    images.use_cache(tmp_path)
    try:
        picture = Image.new("L", (80, 60), 255)
        assert images.examine(picture).size == images.examine(picture).size == (80, 60)
        assert len(calls) == 1 and len(list(tmp_path.glob("*.pkl"))) == 1
    finally:
        images.use_cache(None)


# --- faces and barcodes ----------------------------------------------------


def test_a_block_of_bars_is_found_without_being_decoded():
    pytest.importorskip("cv2")
    card = Image.new("L", (600, 300), 255)
    draw = ImageDraw.Draw(card)
    x, widths = 100, (2, 3, 5, 3, 2, 5)
    while x < 400:
        bar = widths[x % len(widths)]
        draw.rectangle((x, 150, x + bar, 220), fill=0)
        x += bar + widths[(x + 1) % len(widths)]
    draw.text((100, 40), "Jane Sample, staff card, valid to the end of the year", fill=0, font=ImageFont.load_default(size=18))
    ((kind, box),) = regions.find(card)
    assert kind == "barcode" and box[0] < 110 and box[2] > 390 and 140 < box[1] < 160 and 210 < box[3] < 230
    assert regions.find(Image.new("L", (600, 300), 255)) == []


# --- one finding per place -------------------------------------------------


def picture_document():
    """A picture read by two engines, as extraction would leave it."""
    doc = Document("d0c", "card.pdf", "pdf")
    region = doc.add(BlockType.IMAGE_REGION, container=Container.IMAGE, page=1, bbox=BBox(0, 0, 400, 200, Unit.PT),
                     native_ref={"engines": ["tesseract", "doctr"], "trusted": True, "agreement": 0.8,
                                 "cover": [["unconfirmed", 300.0, 150.0, 380.0, 170.0]]})

    def line(text, engine, y, pieces):
        words = [Word(a, b, BBox(x0, y, x1, y + 12, Unit.PT), 0.9) for a, b, x0, x1 in pieces]
        return doc.add(BlockType.PARAGRAPH, text, container=Container.IMAGE, source=Source.OCR, page=1,
                       parent_id=region.block_id, words=words, native_ref={"engine": engine})

    first = line("Born 28 February 1979", "tesseract", 40, [(0, 4, 10, 40), (5, 7, 44, 58), (8, 16, 62, 120), (17, 21, 124, 150)])
    second = line("February 1979", "doctr", 40, [(0, 8, 62, 120), (9, 13, 124, 150)])
    found = [
        make(doc, first, 5, 21, entity="DOB", detector="rules", rule="dob", confidence=0.8, location="image"),
        make(doc, second, 0, 13, entity="DOB", detector="rules", rule="dob", confidence=0.8, location="image"),
    ]
    return doc, region, found


def test_a_value_both_engines_read_is_reported_once_and_both_readings_are_covered():
    doc, region, found = picture_document()
    (kept,) = merge.one_per_place([doc], found)
    assert kept.text == "February 1979" and "two_engines" in kept.flags  # between equals, the neural reading
    # The reading that was dropped covered the day as well: its rectangle is kept for masking.
    assert ["finding", 44, 40, 150, 52] in region.native_ref["cover"]
    assert merge.one_per_place([doc], found) and region.native_ref["cover"].count(["finding", 44, 40, 150, 52]) == 1


def test_findings_in_different_places_or_in_ordinary_text_are_all_kept():
    doc, region, found = picture_document()
    body = doc.add(BlockType.PARAGRAPH, "Born 28 February 1979", page=1)
    plain = make(doc, body, 5, 21, entity="DOB", detector="rules", rule="dob", confidence=0.8, location="body")
    assert len(merge.one_per_place([doc], [found[0], plain])) == 2


# --- what is covered in a picture ------------------------------------------


def test_shapes_covered_inside_pictures_whether_detected_or_not():
    from cerberus.detect.allowlist import shaped_like_pii

    def shapes(text):
        return sorted({text[a:b] for a, b in shaped_like_pii(text)})

    assert shapes("Owner J. Q. Sample signed") == ["J. Q. Sample"]  # the title in front does not hide the name
    assert shapes("Write to j.sample@corp.example today") == ["j.sample@corp.example"]
    assert shapes("Call 555 010 0101 or see page 12") == ["555 010 0101"]
    assert shapes("Staff EMP-00001, born 01/02/1980") == ["01/02/1980", "EMP-00001,"]
    # A short reference code, a year and a count are not personal data.
    assert shapes("REF-0042 raised in 2026 for 1200 items, see section 12.4") == []
    assert {"28 February 1979", "3rd Mar 2020"} <= set(shapes("Born 28 February 1979, joined 3rd Mar 2020"))
    assert shapes("Jane Sample\nOmar Placeholder") == ["Jane Sample", "Omar Placeholder"]
    assert shapes("Executive attestation and Risk Owner") == []


def test_a_trusted_picture_is_covered_at_findings_and_pii_shapes_and_nothing_else():
    doc, region, found = picture_document()
    tokenized, _ = tokens.assign(merge.one_per_place([doc], found))
    boxes = office.picture_boxes(doc, tokenized)[region.block_id]
    # The finding, and the wider reading that was dropped as its duplicate.
    assert {(62, 40, 120, 52), (124, 40, 150, 52), (44, 40, 150, 52)} <= set(boxes)
    # A word the engines merely read differently is not covered, nor is anything else.
    assert all(40 <= box[1] and box[3] <= 52 for box in boxes)
    assert office.picture_boxes(doc, tokenized, pictures="whole") == {region.block_id: None}


def test_ordinary_words_come_from_the_document_itself():
    doc = Document("d0c", "policy.pdf", "pdf")
    doc.add(BlockType.PARAGRAPH, "Payment processing is effective; write to j.sample@corp.example. Jane Sample agreed.")
    box = BBox(0, 0, 1, 1, Unit.PT)
    doc.add(BlockType.PARAGRAPH, "third-party access qzxv", source=Source.OCR,
            words=[Word(0, 11, box, 0.9), Word(12, 18, box, 0.9), Word(19, 23, box, 0.3)])
    picture = doc.add(BlockType.IMAGE_REGION, container=Container.IMAGE)
    doc.add(BlockType.PARAGRAPH, "hidden words", container=Container.IMAGE, parent_id=picture.block_id)
    common = office.common_words(doc)
    # Small-letter words from reliable text; nothing from an address, a name, a picture or an unsure reading.
    assert common == {"processing", "is", "effective", "write", "to", "agreed", "third", "party", "access"}

    assert office._harmless("Effective", common) and office._harmless("processlng,", common)  # one letter misread
    assert office._harmless("REF-0042", common) and office._harmless("2026", common) and office._harmless("ml", common)
    assert not office._harmless("Sample", common) and not office._harmless("jsamplecorpexample", common)
    assert not office._harmless("j.sample@corp.example", common) and not office._harmless("EMP-00001", common)


def test_a_word_the_engines_disagree_on_is_covered_unless_one_read_an_ordinary_word():
    doc = Document("d0c", "form.pdf", "pdf")
    doc.add(BlockType.PARAGRAPH, "the payment was approved after review")

    def box(x):
        return BBox(x, 10, x + 60, 22, Unit.PT)

    doubtful = [["unconfirmed", x, 10, x + 60, 22] for x in (0, 100, 200, 300)]
    region = doc.add(BlockType.IMAGE_REGION, container=Container.IMAGE, page=1,
                     native_ref={"trusted": True, "cover": doubtful})

    def line(text, engine, xs):
        words, at = [], 0
        for piece, x in zip(text.split(), xs):
            words.append(Word(at, at + len(piece), box(x), 0.9))
            at += len(piece) + 1
        doc.add(BlockType.PARAGRAPH, text, container=Container.IMAGE, source=Source.OCR, page=1,
                parent_id=region.block_id, words=words, native_ref={"engine": engine})

    line("paymert zqxwv REF-0042 vkjqz", "tesseract", (0, 100, 200, 300))
    line("payment zqxvw REF-0043 approved", "doctr", (0, 100, 200, 300))
    covered = sorted({b[0] for b in office.picture_boxes(doc, [])[region.block_id]})
    # Only the place where both readings are nonsense. A misread ordinary word, a short
    # code and a place where one engine read an ordinary word are all left visible.
    assert covered == [100]


@pytest.mark.parametrize("change", [{"trusted": False}, {"quarantined": True}])
def test_a_picture_that_is_not_trusted_is_masked_whole(change):
    doc, region, found = picture_document()
    region.quarantined = change.get("quarantined", False)
    region.native_ref["trusted"] = change.get("trusted", True)
    tokenized, _ = tokens.assign(found)
    assert office.picture_boxes(doc, tokenized) == {region.block_id: None}


def test_a_covered_picture_is_rebuilt_from_pixels_with_the_rectangles_painted():
    source = Image.new("RGB", (200, 100), "white")
    ImageDraw.Draw(source).rectangle((50, 40, 120, 60), fill="black")
    blob = io.BytesIO()
    source.save(blob, format="JPEG", comment=b"taken by Jane Sample")
    masked = Image.open(io.BytesIO(office.covered_picture(blob.getvalue(), "image/jpeg", [(50, 40, 120, 60)])))
    assert masked.format == "PNG" and masked.size == (200, 100) and "comment" not in masked.info
    assert masked.getpixel((85, 50)) == (224, 224, 224) and masked.getpixel((10, 10))[0] > 240
    whole = Image.open(io.BytesIO(office.covered_picture(blob.getvalue(), "image/jpeg", None, (200, 100))))
    assert whole.mode == "L" and whole.getpixel((2, 2)) == 224


def card() -> bytes:
    picture = Image.new("RGB", (640, 200), "white")
    draw = ImageDraw.Draw(picture)
    font = ImageFont.load_default(size=28)
    draw.text((30, 30), "User record", fill="black", font=font)
    draw.text((30, 110), "Badge EMP-00002", fill="black", font=font)
    out = io.BytesIO()
    picture.save(out, format="PNG")
    return out.getvalue()


def test_a_trusted_picture_in_a_deck_keeps_its_content_and_loses_the_finding(tmp_path):
    from cerberus.detect import rules

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.shapes.add_picture(io.BytesIO(card()), Emu(500_000), Emu(500_000))
    path = tmp_path / "deck.pptx"
    prs.save(str(path))

    doc = extract_pptx(path)
    (region,) = [b for b in doc.blocks if b.type is BlockType.IMAGE_REGION]
    if not doc.children(region.block_id):
        pytest.skip("Tesseract is not installed")
    # Tests read with one engine; say a second one agreed, as extraction would record it.
    region.native_ref = {**region.native_ref, "trusted": True, "cover": []}
    tokenized, _ = tokens.assign(rules.detect([doc]))
    assert [t.finding.entity for t in tokenized] == ["STAFF_ID"]

    out = mask_pptx(path, doc, tokenized, tmp_path / "masked.pptx")
    (shape,) = Presentation(str(out)).slides[0].shapes
    masked = Image.open(io.BytesIO(shape.image.blob)).convert("L")
    assert masked.size == (640, 200)
    box = doc.get(tokenized[0].finding.block_id).span_boxes(tokenized[0].finding.start, tokenized[0].finding.end)[0]
    inside = masked.crop((int(box.x0), int(box.y0), int(box.x1), int(box.y1)))
    assert inside.getextrema() == (224, 224)  # the staff number is painted over
    assert masked.crop((30, 30, 200, 60)).getextrema()[0] < 80  # "User record" is still there
    assert verify.check_file(out, tokenized, rules.detect) == []


# --- falling back -----------------------------------------------------------


def test_a_picture_that_still_shows_pii_after_masking_stops_being_trusted():
    doc, region, _ = picture_document()
    other = doc.add(BlockType.IMAGE_REGION, container=Container.IMAGE, page=2, native_ref={"trusted": True})
    in_text = verify.Problem("card.pdf", "rescan", "page 1, body", "PERSON", 1, False)
    assert not pipeline.distrust_pictures(doc, [in_text]) and region.native_ref["trusted"]
    in_picture = verify.Problem("card.pdf", "rescan", "page 1, image", "PERSON", 1, True)
    assert pipeline.distrust_pictures(doc, [in_picture])
    assert not region.native_ref["trusted"] and "fallback" in region.native_ref
    assert other.native_ref["trusted"]  # a picture on another page of a PDF is left alone
    assert not pipeline.distrust_pictures(doc, [in_picture])  # nothing left to change


def test_in_an_office_file_only_the_picture_that_failed_stops_being_trusted():
    deck = Document("d3ck", "deck.pptx", "pptx")
    first, second, third = (
        deck.add(BlockType.IMAGE_REGION, container=Container.IMAGE, page=1, native_ref={"trusted": True}) for _ in range(3)
    )
    problem = verify.Problem("deck.pptx", "rescan", "page 1, image", "PERSON", 1, True, picture=(1, 3))
    assert pipeline.distrust_pictures(deck, [problem])
    assert [b.native_ref["trusted"] for b in (first, second, third)] == [True, False, True]
    # If the masked file does not have the same pictures, there is no telling which one: all of them.
    unclear = verify.Problem("deck.pptx", "rescan", "page 1, image", "PERSON", 1, True, picture=(0, 5))
    assert pipeline.distrust_pictures(deck, [unclear])
    assert [b.native_ref["trusted"] for b in (first, second, third)] == [False, False, False]


# --- the real engine --------------------------------------------------------


@pytest.mark.neural
@pytest.mark.skipif(not neural.available(), reason="docTR or its weights are not installed")
def test_two_engines_read_a_clear_picture_and_agree():
    picture = Image.open(io.BytesIO(card()))
    reading = images._examine(picture.convert("L"), 0, picture)
    assert reading.second is not None and reading.trusted and reading.agreement > 0.8
    text = " ".join(w.text for line in reading.second for w in line.words)
    assert "EMP-00002" in text and "Badge" in text
