import io

import docx
import pytest
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml import parse_xml
from PIL import Image

from cerberus.extract.docx import extract_docx
from cerberus.extract.ir import BlockType, Container

# All names and contact details below are invented.

W = 'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"'
R = 'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"'
MC = 'xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006"'


@pytest.fixture
def document(tmp_path):
    d = docx.Document()
    d.core_properties.author = "Jane Sample"

    d.add_heading("1. Contacts", level=1)
    d.add_paragraph("Reach the helpdesk by email.")
    d.add_paragraph("")  # empty paragraphs are spacing, not content
    d.add_heading("1.1 Owners", level=2)
    d.add_paragraph("Jane Sample owns onboarding.", style="List Bullet")

    table = d.add_table(rows=3, cols=2)
    for r, row in enumerate([("Name", "Mobile"), ("Jane Sample", "+1 (555) 0101"),
                             ("Omar Placeholder", "")]):
        for c, value in enumerate(row):
            table.cell(r, c).text = value
    table.cell(1, 1).merge(table.cell(2, 1))  # one mobile number spanning two rows
    table.cell(0, 0).add_paragraph("(full name)")

    d.add_heading("2. Evidence", level=1)
    png = io.BytesIO()
    Image.new("RGB", (40, 20), "white").save(png, format="PNG")
    png.seek(0)
    d.add_picture(png)

    # A link whose address is not in the visible text.
    link = d.add_paragraph("Contact: ")
    rid = d.part.relate_to("mailto:j.sample@corp.example", RT.HYPERLINK, is_external=True)
    link._p.append(parse_xml(f'<w:hyperlink {W} {R} r:id="{rid}"><w:r><w:t>Jane</w:t></w:r></w:hyperlink>'))

    # A tracked deletion: gone from the page, still in the file.
    tracked = d.add_paragraph("Approved by the board")
    tracked._p.append(parse_xml(
        f'<w:del {W} w:author="Omar Placeholder"><w:r><w:delText> and Jane Sample</w:delText></w:r></w:del>'
    ))

    # A text box, stored twice (for new and old Word) as Word really does.
    box = "<w:txbxContent><w:p><w:r><w:t>Boxed note</w:t></w:r></w:p></w:txbxContent>"
    d.add_paragraph("See box.")._p.append(parse_xml(
        f"<w:r {W} {MC}><mc:AlternateContent><mc:Choice Requires=\"wps\"><w:drawing>{box}</w:drawing>"
        f"</mc:Choice><mc:Fallback><w:pict>{box}</w:pict></mc:Fallback></mc:AlternateContent></w:r>"
    ))

    # Soft line break and tab inside one paragraph.
    run = d.add_paragraph().add_run("Line one")
    run.add_break()
    run.add_text("Line\ttwo")

    d.sections[0].header.paragraphs[0].text = "Internal use"
    d.sections[0].footer.paragraphs[0].text = "Prepared by Jane Sample"

    path = tmp_path / "doc.docx"
    d.save(str(path))
    return extract_docx(path)


def of(doc, type_):
    return [b for b in doc.blocks if b.type is type_]


def find(doc, text):
    (block,) = [b for b in doc.blocks if b.text == text]
    return block


def test_valid_and_has_no_pages(document):
    assert document.validate() == []
    assert document.pages == []
    assert all(b.page is None and b.bbox is None for b in document.blocks)


def test_headings_nest_by_level_and_set_section_path(document):
    h1, h11, h2 = of(document, BlockType.HEADING)
    assert (h1.parent_id, h11.parent_id, h2.parent_id) == (None, h1.block_id, None)
    assert h11.section_path == ("1. Contacts",)
    item = find(document, "Jane Sample owns onboarding.")
    assert item.type is BlockType.LIST_ITEM
    assert item.parent_id == h11.block_id
    assert item.section_path == ("1. Contacts", "1.1 Owners")
    assert find(document, "Reach the helpdesk by email.").type is BlockType.PARAGRAPH
    assert not [b for b in document.blocks if b.type is BlockType.PARAGRAPH and not b.text.strip()]


def test_table_is_in_document_order_with_headers_and_merges(document):
    (table,) = of(document, BlockType.TABLE)
    assert find(document, "Jane Sample owns onboarding.").order < table.order
    assert table.order < find(document, "2. Evidence").order
    cells = document.children(table.block_id)
    # The merged cell's lower half is not emitted; a two-paragraph cell is one block.
    assert [c.text for c in cells] == [
        "Name\n(full name)", "Mobile", "Jane Sample", "+1 (555) 0101", "Omar Placeholder",
    ]
    assert [(c.row, c.col) for c in cells] == [(0, 0), (0, 1), (1, 0), (1, 1), (2, 0)]
    assert document.get(cells[3].header_id).text == "Mobile"
    assert all(c.container is Container.TABLE for c in cells)


def test_image_is_placed_where_it_is_anchored(document):
    (image,) = of(document, BlockType.IMAGE_REGION)
    assert image.parent_id == find(document, "2. Evidence").block_id
    assert image.native_ref["content_type"] == "image/png"
    assert image.native_ref["target"].startswith("/word/media/")
    assert len(image.native_ref["sha1"]) == 40


def test_text_not_visible_on_the_page_is_emitted(document):
    blocks = of(document, BlockType.METADATA)
    # python-docx's default template fills in a few file properties of its own.
    assert ("author", "Jane Sample") in [(b.native_ref.get("core_property"), b.text) for b in blocks]
    meta = {next(k for k in b.native_ref if k not in ("part", "p")): b for b in blocks}
    assert meta["hyperlink"].text == "mailto:j.sample@corp.example"
    assert meta["hyperlink"].parent_id == find(document, "Contact: Jane").block_id
    assert meta["tracked_deletion"].text == " and Jane Sample"
    assert meta["revision_author"].text == "Omar Placeholder"
    assert find(document, "Approved by the board")  # deleted words are not in the visible text


def test_text_box_is_read_once(document):
    (box,) = of(document, BlockType.TEXT_BOX)
    assert box.text == "Boxed note"
    assert box.parent_id == find(document, "See box.").block_id


def test_breaks_and_tabs_are_kept_as_characters(document):
    assert find(document, "Line one\nLine\ttwo")


def test_headers_and_footers(document):
    (header,) = of(document, BlockType.HEADER)
    (footer,) = of(document, BlockType.FOOTER)
    assert [b.text for b in document.children(header.block_id)] == ["Internal use"]
    assert [b.text for b in document.children(footer.block_id)] == ["Prepared by Jane Sample"]
