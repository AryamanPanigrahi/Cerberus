import io
import zipfile

import docx
import pytest
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml import parse_xml
from PIL import Image
from pptx import Presentation
from pptx.util import Emu, Pt

from cerberus.detect import rules
from cerberus.detect.finding import make
from cerberus.extract.docx import extract_docx
from cerberus.extract.ir import BlockType
from cerberus.extract.pptx import extract_pptx
from cerberus.redact import office, tokens
from cerberus.redact.docx import mask_docx
from cerberus.redact.pptx import mask_pptx

# All names and contact details below are invented.

NAMES = ("Jane Sample", "Omar Placeholder")
SECRETS = ("Jane", "Sample", "Omar", "Placeholder", "j.sample", "010-0101", "EMP-00001")
# For searching a file's raw XML, whole values and distinctive parts only: a surname that is also
# an English word collides with Word's own built-in style names ("Placeholder Text").
RAW_SECRETS = ("Jane Sample", "Omar Placeholder", "Jane", "Omar", "j.sample", "010-0101", "EMP-00001")
W = 'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"'
R = 'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"'
MC = 'xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006"'


def picture() -> io.BytesIO:
    out = io.BytesIO()
    Image.new("RGB", (30, 20), "white").save(out, format="PNG")  # too small to be OCR'd
    out.seek(0)
    return out


def findings_for(doc):
    """Rule findings plus every occurrence of the known names, as the name detector would give."""
    found = rules.detect([doc])
    for block in doc.blocks:
        for name in NAMES:
            start = block.text.find(name)
            while start >= 0:
                found.append(make(doc, block, start, start + len(name), entity="PERSON", detector="test",
                                  rule="t", confidence=0.95, location="body"))
                start = block.text.find(name, start + 1)
    return tokens.assign(found)[0]


def raw_text(path) -> str:
    """Everything in every XML part of the file: what a search of the raw contents would see."""
    with zipfile.ZipFile(path) as archive:
        return " ".join(archive.read(n).decode("utf-8", "ignore") for n in archive.namelist()
                        if n.endswith((".xml", ".rels")))


# --- the shared run logic -------------------------------------------------


def test_token_goes_into_the_first_run_and_other_runs_keep_their_text():
    pieces = ["Signed by Ja", "ne Sam", "ple today, and ", "Omar", "."]
    spans = [(10, 21, "<PERSON_1>"), (33, 37, "<PERSON_2>")]
    assert office.apply_spans(pieces, spans) == ["Signed by <PERSON_1>", "", " today, and ", "<PERSON_2>", "."]
    assert office.apply_spans(["no change"], []) == ["no change"]


def test_a_span_outside_the_text_is_refused():
    with pytest.raises(ValueError):
        office.apply_spans(["short"], [(10, 14, "<X_1>")])


# --- DOCX ---------------------------------------------------------------


@pytest.fixture
def word_file(tmp_path):
    d = docx.Document()
    d.core_properties.author = "Jane Sample"
    d.core_properties.last_modified_by = "Somebody Unseen"  # a name no detector reports
    d.add_heading("1. Approvals", level=1)
    para = d.add_paragraph("Approved by ")
    para.add_run("Jane ").bold = True  # the name straddles two runs
    para.add_run("Sample (EMP-00001), reachable at j.sample@corp.example.")
    table = d.add_table(rows=2, cols=2)
    for r, row in enumerate([("NAME", "MOBILE"), ("Omar Placeholder", "+1 (555) 010-0101")]):
        for c, value in enumerate(row):
            table.cell(r, c).text = value
    d.add_picture(picture())
    link = d.add_paragraph("Write: ")
    rid = d.part.relate_to("mailto:j.sample@corp.example", RT.HYPERLINK, is_external=True)
    link._p.append(parse_xml(f'<w:hyperlink {W} {R} r:id="{rid}"><w:r><w:t>here</w:t></w:r></w:hyperlink>'))
    tracked = d.add_paragraph("Signed off")
    tracked._p.append(parse_xml(
        f'<w:del {W} w:author="Omar Placeholder"><w:r><w:delText> by Jane Sample</w:delText></w:r></w:del>'))
    box = "<w:txbxContent><w:p><w:r><w:t>Note for Jane Sample</w:t></w:r></w:p></w:txbxContent>"
    d.add_paragraph("See box.")._p.append(parse_xml(
        f"<w:r {W} {MC}><mc:AlternateContent><mc:Choice Requires=\"wps\"><w:drawing>{box}</w:drawing>"
        f"</mc:Choice><mc:Fallback><w:pict>{box}</w:pict></mc:Fallback></mc:AlternateContent></w:r>"))
    d.sections[0].footer.paragraphs[0].text = "Prepared by Jane Sample"
    path = tmp_path / "doc.docx"
    d.save(str(path))
    return path


def test_masked_docx_holds_tokens_and_none_of_the_original_values(word_file, tmp_path):
    doc = extract_docx(word_file)
    out = mask_docx(word_file, doc, findings_for(doc), tmp_path / "masked" / "doc.docx")

    again = extract_docx(out)
    text = " ".join(b.text for b in again.blocks)
    for secret in SECRETS + ("Unseen",):
        assert secret not in text
    assert "Approved by <PERSON_1> (<STAFF_ID_1>), reachable at <EMAIL_1>." in text
    assert "Note for <PERSON_1>" in text and "Prepared by <PERSON_1>" in text
    assert "<PERSON_2>" in text and "<PHONE_1>" in text

    # Nothing is left in any part of the file, including the copies the extractor does not read.
    raw = raw_text(out)
    assert [secret for secret in RAW_SECRETS + ("Unseen",) if secret in raw] == []
    # The text box is still stored twice (current Word and Word 2007), and both copies are masked.
    assert "mc:Fallback" in raw and raw.count("Note for &lt;PERSON_1&gt;") == 2


def test_masked_docx_keeps_its_structure_and_formatting(word_file, tmp_path):
    doc = extract_docx(word_file)
    out = mask_docx(word_file, doc, findings_for(doc), tmp_path / "masked.docx")
    again = extract_docx(out)
    kinds = lambda d: [b.type for b in d.blocks if b.type not in (BlockType.METADATA, BlockType.PARAGRAPH)]  # noqa: E731
    assert kinds(again) == kinds(doc)
    bold = [r.text for p in docx.Document(str(out)).paragraphs for r in p.runs if r.bold]
    assert bold == ["<PERSON_1>"]  # the bold run still exists and now holds the token


def test_masked_docx_replaces_pictures_and_rewrites_properties(word_file, tmp_path):
    doc = extract_docx(word_file)
    out = mask_docx(word_file, doc, findings_for(doc), tmp_path / "masked.docx")
    original = [b.native_ref["sha1"] for b in doc.blocks if b.type is BlockType.IMAGE_REGION]
    masked = [b.native_ref["sha1"] for b in extract_docx(out).blocks if b.type is BlockType.IMAGE_REGION]
    assert len(masked) == len(original) == 1 and masked != original
    with zipfile.ZipFile(out) as archive:
        media = [n for n in archive.namelist() if n.startswith("word/media/")]
    assert len(media) == 1  # the original image file is gone, not just unreferenced

    properties = docx.Document(str(out)).core_properties
    assert properties.author == "<PERSON_1>"
    assert properties.last_modified_by == ""  # nothing was detected in it, so it is emptied


def test_a_document_that_changed_since_it_was_scanned_is_refused(word_file, tmp_path):
    doc = extract_docx(word_file)
    tokenized = findings_for(doc)
    changed = docx.Document(str(word_file))
    changed.paragraphs[1].runs[0].text = "Rejected by "
    other = tmp_path / "changed.docx"
    changed.save(str(other))
    with pytest.raises(ValueError, match="no longer matches"):
        mask_docx(other, doc, tokenized, tmp_path / "masked.docx")


# --- PPTX ---------------------------------------------------------------


@pytest.fixture
def deck_file(tmp_path):
    prs = Presentation()
    prs.core_properties.author = "Jane Sample"
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    title = slide.shapes.add_textbox(Emu(500_000), Emu(200_000), Emu(6_000_000), Emu(500_000))
    title.text_frame.text = "Team"
    title.text_frame.paragraphs[0].runs[0].font.size = Pt(28)
    card = slide.shapes.add_textbox(Emu(500_000), Emu(1_500_000), Emu(4_000_000), Emu(900_000))
    first = card.text_frame.paragraphs[0]
    first.add_run().text = "Jane "
    first.add_run().text = "Sample"
    first.runs[0].font.bold = True
    card.text_frame.add_paragraph().text = "j.sample@corp.example"
    card.text_frame.add_paragraph().text = "M +1 (555) 010-0101"
    frame = slide.shapes.add_table(2, 2, Emu(500_000), Emu(3_000_000), Emu(4_000_000), Emu(600_000))
    for r, row in enumerate([("NAME", "ROLE"), ("Omar Placeholder", "Lead")]):
        for c, value in enumerate(row):
            frame.table.cell(r, c).text = value
    slide.shapes.add_picture(picture(), Emu(6_000_000), Emu(1_500_000), Emu(600_000), Emu(400_000))
    slide.notes_slide.notes_text_frame.text = "Ask Jane Sample first."
    path = tmp_path / "deck.pptx"
    prs.save(str(path))
    return path


def test_masked_pptx_holds_tokens_and_none_of_the_original_values(deck_file, tmp_path):
    doc = extract_pptx(deck_file)
    out = mask_pptx(deck_file, doc, findings_for(doc), tmp_path / "masked" / "deck.pptx")

    again = extract_pptx(out)
    text = " | ".join(b.text for b in again.blocks)
    for secret in SECRETS:
        assert secret not in text
    assert "<PERSON_1>\n<EMAIL_1>\nM <PHONE_1>" in text
    assert "<PERSON_2>" in text and "Ask <PERSON_1> first." in text
    raw = raw_text(out)
    assert [secret for secret in RAW_SECRETS if secret in raw] == []


def test_masked_pptx_keeps_layout_formatting_and_replaces_pictures(deck_file, tmp_path):
    doc = extract_pptx(deck_file)
    out = mask_pptx(deck_file, doc, findings_for(doc), tmp_path / "masked.pptx")
    again = extract_pptx(out)
    shapes = lambda d: [(b.type, b.bbox) for b in d.blocks if b.type is not BlockType.METADATA]  # noqa: E731
    assert shapes(again) == shapes(doc)  # every shape is where it was

    prs = Presentation(str(out))
    card = [s for s in prs.slides[0].shapes if s.has_text_frame][1]
    runs = card.text_frame.paragraphs[0].runs
    assert [(r.text, r.font.bold) for r in runs] == [("<PERSON_1>", True), ("", None)]
    assert prs.core_properties.author == "<PERSON_1>"
    with zipfile.ZipFile(out) as archive:
        media = [n for n in archive.namelist() if n.startswith("ppt/media/")]
    assert len(media) == 1
    original = [b.native_ref["sha1"] for b in doc.blocks if b.type is BlockType.IMAGE_REGION]
    masked = [b.native_ref["sha1"] for b in again.blocks if b.type is BlockType.IMAGE_REGION]
    assert masked != original


def test_an_old_word_copy_that_cannot_be_matched_is_removed(tmp_path):
    d = docx.Document()
    box = "<w:txbxContent><w:p><w:r><w:t>Note for Jane Sample</w:t></w:r></w:p></w:txbxContent>"
    art = '<v:shape xmlns:v="urn:schemas-microsoft-com:vml"><v:textpath string="Jane Sample"/></v:shape>'
    d.add_paragraph("See box.")._p.append(parse_xml(
        f"<w:r {W} {MC}><mc:AlternateContent><mc:Choice Requires=\"wps\"><w:drawing>{box}</w:drawing>"
        f"</mc:Choice><mc:Fallback><w:pict>{box}{art}</w:pict></mc:Fallback></mc:AlternateContent></w:r>"))
    path = tmp_path / "art.docx"
    d.save(str(path))
    doc = extract_docx(path)
    raw = raw_text(mask_docx(path, doc, findings_for(doc), tmp_path / "masked.docx"))
    assert "Jane" not in raw and "mc:Fallback" not in raw
    assert "Note for &lt;PERSON_1&gt;" in raw


def test_a_span_over_a_line_break_leaves_the_break_and_clears_both_sides():
    pieces = ["Jane Sample", "\n", "Omar Placeholder"]
    out = office.apply_spans(pieces, [(0, 28, "<PERSON_1>")], fixed=frozenset({1}))
    assert out == ["<PERSON_1>", "\n", ""]
