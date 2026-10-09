import shutil
import zipfile

import docx
import pymupdf
import pytest
from pptx import Presentation
from pptx.util import Emu

from cerberus import pipeline
from cerberus.detect import rules
from cerberus.detect.finding import Finding
from cerberus.extract.docx import extract_docx
from cerberus.extract.ir import BlockType, Document
from cerberus.redact import tokens, verify
from cerberus.redact.docx import mask_docx
from cerberus.redact.text import redacted_text

# All names and contact details below are invented.


def word_file(folder):
    d = docx.Document()
    d.add_heading("1. Contacts", level=1)
    para = d.add_paragraph("Write to ")
    para.add_run("j.sample").bold = True  # the address straddles two runs
    para.add_run("@corp.example or call +1 (555) 010-0101.")
    d.add_paragraph("Staff number EMP-00001 is on file.")
    path = folder / "contacts.docx"
    d.save(str(path))
    return path


def deck_file(folder):
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    box = slide.shapes.add_textbox(Emu(500_000), Emu(500_000), Emu(5_000_000), Emu(600_000))
    box.text_frame.text = "Helpdesk: helpdesk@corp.example"
    path = folder / "deck.pptx"
    prs.save(str(path))
    return path


def typed_pdf(folder):
    pdf = pymupdf.open()
    page = pdf.new_page()
    page.insert_text((72, 100), "Escalations go to duty.manager@corp.example within the hour, without", fontsize=11)
    page.insert_text((72, 116), "exception, and are logged by the service desk the same day.", fontsize=11)
    pdf.set_metadata({"author": "duty.manager@corp.example"})
    path = folder / "policy.pdf"
    pdf.save(str(path))
    return path


def scanned(path):
    doc = extract_docx(path)
    tokenized, _ = tokens.assign(rules.detect([doc]))
    return doc, tokenized


def test_a_properly_masked_file_passes_both_checks(tmp_path):
    source = word_file(tmp_path)
    doc, tokenized = scanned(source)
    out = mask_docx(source, doc, tokenized, tmp_path / "masked.docx")
    assert verify.check_file(out, tokenized, rules.detect) == []
    text = redacted_text(doc, tokenized)
    assert verify.check_text("contacts.md", text, tokenized, rules.detect) == []
    verify.ensure_clean([])  # nothing to report: no error


def test_an_unmasked_file_fails_both_checks_without_quoting_the_values(tmp_path):
    source = word_file(tmp_path)
    doc, tokenized = scanned(source)
    problems = verify.check_file(source, tokenized, rules.detect)  # the original, passed off as masked
    assert {p.check for p in problems} == {"raw", "rescan"}
    assert {p.what for p in problems if p.check == "rescan"} == {"EMAIL", "PHONE", "STAFF_ID"}
    with pytest.raises(verify.VerificationError) as error:
        verify.ensure_clean(problems)
    message = str(error.value)
    assert "nothing was released" in message
    assert "j.sample" not in message and "010-0101" not in message  # the report can be shown and logged


def test_raw_search_finds_a_value_the_extractor_does_not_read(tmp_path):
    source = word_file(tmp_path)
    doc, tokenized = scanned(source)
    out = mask_docx(source, doc, tokenized, tmp_path / "masked.docx")
    # Plant the address in a part the extractor never opens.
    tampered = tmp_path / "tampered.docx"
    with zipfile.ZipFile(out) as original, zipfile.ZipFile(tampered, "w") as copy:
        for item in original.infolist():
            copy.writestr(item, original.read(item.filename))
        copy.writestr("customXml/item9.xml", "<note>sent by <b>j.sample</b>@corp.example</note>")
    problems = verify.check_file(tampered, tokenized, rules.detect)
    assert [(p.check, p.where, p.what) for p in problems] == [("raw", "customXml/item9.xml", "<EMAIL_1>")]


def test_values_split_by_markup_or_line_breaks_are_still_found():
    values = {"j.sample@corp.example": "<EMAIL_1>", "jane sample": "<PERSON_1>"}
    assert verify.find_values("mail J.Sample@Corp.Example now", values) == ["<EMAIL_1>"]
    assert verify.find_values("Jane\n   Sample signed", values) == ["<PERSON_1>"]
    assert verify.find_values("JaneSample", values) == ["<PERSON_1>"]  # as when two runs are joined
    assert verify.find_values("<PERSON_1> signed", values) == []


def test_short_values_are_left_to_the_rescan():
    # Four digits would match stray numbers in a file's own markup.
    found = [Finding("d", "b", 0, len(v), v, "SSN", "rules", "r", 0.9) for v in ("3456", "900-12-3456")]
    tokenized, _ = tokens.assign(found)
    assert list(verify.known_values(tokenized)) == ["900-12-3456"]


def test_tokens_and_notices_in_the_output_are_not_reported(tmp_path):
    text = "(picture withheld: 2 items of PII detected inside)\nContact <EMAIL_1> or <PHONE_1>.\n"
    assert verify.check_text("out.md", text, [], rules.detect) == []


# --- the whole pipeline ---------------------------------------------------


def test_pipeline_releases_verified_outputs_for_all_three_formats(tmp_path):
    source = tmp_path / "in"
    source.mkdir()
    paths = [word_file(source), deck_file(source), typed_pdf(source)]
    result = pipeline.run(sorted(paths), tmp_path / "out", use_ner=False)

    assert result.problems == []
    assert sorted(p.name for p in result.masked) == ["contacts.docx", "deck.pptx", "policy.pdf"]
    assert sorted(p.name for p in result.redacted) == ["contacts.md", "deck.md", "policy.md"]
    assert not (tmp_path / "out" / "pending").exists() and not (tmp_path / "out" / "withheld").exists()
    text = (tmp_path / "out" / "redacted" / "contacts.md").read_text(encoding="utf-8")
    assert "Write to <EMAIL_" in text and "j.sample" not in text
    with pymupdf.open(str(tmp_path / "out" / "masked" / "policy.pdf")) as masked:
        assert "duty.manager" not in masked[0].get_text() and not (masked.metadata or {}).get("author")
    assert (tmp_path / "out" / "vault.json").exists()
    assert [p.name for p in result.reports] == [
        "audit.csv", "contacts.summary.md", "deck.summary.md", "policy.summary.md", "exposure.md",
    ]
    report = "".join(p.read_text(encoding="utf-8-sig") for p in result.reports)
    assert "**passed**" in report and "<EMAIL_" in report
    assert [value for value in ("j.sample", "010-0101", "helpdesk@", "duty.manager") if value in report] == []


def test_pipeline_withholds_everything_when_an_output_is_not_clean(tmp_path, monkeypatch):
    source = tmp_path / "in"
    source.mkdir()
    paths = [word_file(source), deck_file(source)]

    def broken(path, doc, tokenized, out, pictures="findings"):  # a masker that masks nothing
        out.parent.mkdir(parents=True, exist_ok=True)
        shutil.copy(path, out)
        return out

    monkeypatch.setitem(pipeline._MASKERS, ".pptx", broken)
    with pytest.raises(verify.VerificationError):
        pipeline.run(sorted(paths), tmp_path / "out", use_ner=False)

    out = tmp_path / "out"
    assert not (out / "masked").exists() and not (out / "redacted").exists() and not (out / "pending").exists()
    # Even the outputs that were fine are held back: a run is released whole or not at all.
    assert (out / "withheld" / "masked" / "contacts.docx").exists()
    assert (out / "withheld" / "masked" / "deck.pptx").exists()
    # The audit report is still written, and says so.
    assert "FAILED" in (out / "report" / "exposure.md").read_text(encoding="utf-8")


def test_looking_again_finds_what_the_first_pass_missed_and_places_it_in_the_original():
    doc = Document("d0c", "sample.docx", "docx")
    block = doc.add(BlockType.TABLE_CELL, "j.sample@corp.example\nOmar Placeholder\no.other@corp.example")
    first = rules.detect([doc])  # the two addresses; no name detector is running
    assert [f.entity for f in first] == ["EMAIL", "EMAIL"]

    def detect(docs):
        """A detector that only sees the name once its neighbours have become tokens."""
        found = rules.detect(docs)
        for d in docs:
            for b in d.blocks:
                if "<EMAIL_1>" in b.text and "Omar Placeholder" in b.text:
                    start = b.text.index("Omar Placeholder")
                    found.append(Finding(d.doc_id, b.block_id, start, start + 16, "Omar Placeholder", "PERSON",
                                         "test", "t", 0.5))
        return found

    found = pipeline.look_again([doc], first, detect)
    (name,) = [f for f in found if f.entity == "PERSON"]
    assert block.text[name.start:name.end] == "Omar Placeholder"  # mapped back past the token before it
    assert "second_look" in name.flags
    assert len(found) == 3
    # Looking again when there is nothing new changes nothing.
    assert pipeline.look_again([doc], found, detect) == found
